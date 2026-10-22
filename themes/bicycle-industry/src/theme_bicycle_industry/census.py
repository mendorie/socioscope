"""経済センサス-活動調査（令和 3 年）卸売業・小売業 産業編（総括表）第 1 表 xlsx → long 行。

e-Stat `file-download?statInfId=000040032840&fileKind=0`（appId 不要、149 KB、1 シート `f1_001`）。
産業分類細分類別・年次別（1994・1997・1999・2002・2004・2007・2012・2014・2016・2021）の
事業所数（合計／法人／個人）、従業者数、年間商品販売額（百万円）、売場面積（㎡）。全国 ＋ 都道府県。
2007 年以前と 2014 年は商業統計、2012・2016・2021 年は経済センサス。2021 年の個人経営は
年間商品販売額・売場面積を含まない（注 7）→ 2016→2021 の販売額は法人中心の比較になる。
対象業種は `INDUSTRIES`（小売業計、自動車小売業、二輪自動車小売業、自転車小売業、機械器具小売業）。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import openpyxl

SOURCE = "estat_census"
LICENSE = (
    "政府標準利用規約（第2.0版）＝CC BY 4.0 互換。出典: 政府統計の総合窓口(e-Stat)、"
    "令和3年経済センサス‐活動調査（総務省・経済産業省）卸売業，小売業 産業編（総括表）第1表 を加工"
    "（確認 2026-10-22）"
)
STAT_INF_ID = "000040032840"
URL = f"https://www.e-stat.go.jp/stat-search/file-download?statInfId={STAT_INF_ID}&fileKind=0"
RAW_NAME = "census_r3_retail_table1.xlsx"
SHEET = "f1_001"
AREA = "00_全国"

# 産業ラベル（ファイル表記そのまま）→ キー
INDUSTRIES: dict[str, str] = {
    "I2_小売業計": "retail_total",
    "591_自動車小売業": "auto_retail",
    "5914_二輪自動車小売業(原動機付自転車を含む)": "motorcycle_retail",
    "592_自転車小売業": "bicycle_retail",
    "593_機械器具小売業(自動車，自転車を除く)": "machinery_retail_other",
}
METRICS: tuple[str, ...] = (
    "establishments",
    "establishments_corporate",
    "establishments_individual",
    "persons_engaged",
    "annual_sales_mn_jpy",
    "sales_floor_m2",
)
HEADER: tuple[str, ...] = (
    "事業所数",
    "事業所数（法人）",
    "事業所数（個人）",
    "従業者数_合計",
    "年間商品販売額",
    "売場面積",
)
# 調査（注 5）: 2007 年以前と 2014 年は商業統計、それ以外は経済センサス
SURVEY_OF_YEAR: dict[int, str] = {
    1994: "commercial_census",
    1997: "commercial_census",
    1999: "commercial_census_simplified",
    2002: "commercial_census",
    2004: "commercial_census_simplified",
    2007: "commercial_census",
    2012: "economic_census",
    2014: "commercial_census",
    2016: "economic_census",
    2021: "economic_census",
}
_YEAR = re.compile(r"^\d+_(\d{4})年$")


@dataclass
class Parsed:
    rows: list[dict[str, object]]
    skipped: list[str] = field(default_factory=list)


def _cell(v: object) -> float | None:
    """数値、"-"（該当なし）→ None、"x"（秘匿）→ None。他の文字列は ValueError。"""
    if v is None:
        return None
    if isinstance(v, int | float):
        return float(v)
    s = str(v).strip()
    if s in {"-", "x", "X", "…", ""}:
        return None
    try:
        return float(s.replace(",", ""))
    except ValueError as e:
        msg = f"non-numeric cell {v!r}"
        raise ValueError(msg) from e


def parse(payload: bytes) -> Parsed:
    """xlsx → long 行 (industry_label, industry, year, survey, metric, value, source)。

    対象業種が 1 つでも欠けると ValueError（様式変更の検知）。全国行だけを使う。
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as e:
        msg = "payload is not an xlsx workbook"
        raise ValueError(msg) from e
    if SHEET not in wb.sheetnames:
        msg = f"sheet {SHEET!r} missing; sheets = {wb.sheetnames}"
        raise ValueError(msg)
    rows = list(wb[SHEET].iter_rows(values_only=True))
    header_idx = next(
        (i for i, r in enumerate(rows) if tuple(str(c) for c in r[4:10]) == HEADER), None
    )
    if header_idx is None:
        msg = "metric header row (事業所数 … 売場面積) not found"
        raise ValueError(msg)
    out: list[dict[str, object]] = []
    skipped: list[str] = []
    seen: set[str] = set()
    for r in rows[header_idx + 1 :]:
        if not r or r[0] != AREA:
            continue
        label = str(r[2])
        if label not in INDUSTRIES:
            continue
        m = _YEAR.match(str(r[3]))
        if not m:
            skipped.append(f"{label}: unrecognised year label {r[3]!r}")
            continue
        year = int(m.group(1))
        seen.add(label)
        for metric, cell in zip(METRICS, r[4:10], strict=True):
            out.append(
                {
                    "industry_label": label,
                    "industry": INDUSTRIES[label],
                    "year": year,
                    "survey": SURVEY_OF_YEAR.get(year, "unknown"),
                    "metric": metric,
                    "value": _cell(cell),
                    "source": SOURCE,
                }
            )
    missing = set(INDUSTRIES) - seen
    if missing:
        msg = f"industries missing from table: {sorted(missing)}"
        raise ValueError(msg)
    return Parsed(rows=out, skipped=skipped)


MART_KEY: tuple[str, ...] = ("industry", "year", "metric")


def build_mart(staged: list[dict[str, object]]) -> list[dict[str, object]]:
    out = [{"iso3": "JPN", **r} for r in staged]
    out.sort(key=lambda r: tuple(str(r[c]) for c in MART_KEY))
    seen: set[tuple[object, ...]] = set()
    for r in out:
        k = tuple(r[c] for c in MART_KEY)
        if k in seen:
            msg = f"duplicate census key {k}"
            raise ValueError(msg)
        seen.add(k)
    return out
