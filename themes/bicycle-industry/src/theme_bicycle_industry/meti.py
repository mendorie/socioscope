"""経済産業省 生産動態統計（機械統計）年報 xlsx → 行 の純粋変換（決定的・fail-closed）。

対象: 調査票 2430「自転車及び車いす」（軽快車／電動アシスト車／その他／車いす）と
2400「自動車」（軽・小型・普通乗用車、排気量別二輪）。同じ調査・同じファイルなので三産業を
そのまま比べられる（design/themes/bicycle-industry.md H2・H6）。

ファイルは「実数表」シートに全品目が入る。1 行目に時間軸コード（YYYY000000 = 暦年、YYYY100000 =
年度、YYYY00ssee = 四半期 ss–ee 月、YYYY00MMMM = 月）、2 行目に列名、3 行目以降が
調査票番号 / 品目番号 / アイテム記号 / 調査票名 / 項目名 / 品目名 / アイテム名 / 単位名 /
値… / 英語名 5 列。
2021–2023 年報は暦年 5 年分・年度 2 年分・四半期 4 本・月次 12 本、2024 年報以降は月次 12 本のみ
（中間計の掲載廃止）。秘匿は "X"、不詳は "***"。

品目番号は年で振り直される（2024 年にトラックの番号が変わった）ので、
品目は**正規化した品目名**で同定する。
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field

import openpyxl

SOURCE = "meti_seidou"
LICENSE = (
    "経済産業省 利用規約（政府標準利用規約 第2.0版 準拠、出典表示で商用含め利用可。"
    "確認 2026-10-22）。出典: 経済産業省「経済産業省生産動態統計年報 機械統計編」"
)
BASE = "https://www.meti.go.jp/statistics/tyo/seidou/result/ichiran/resourceData/08_seidou/nenpo"
# 年報 xlsx が同じ URL 形式で置かれている年。2007–2020 は業種別 .xls（別形式、未対応）。
YEARS: tuple[int, ...] = (2021, 2022, 2023, 2024, 2025)
SHEET = "実数表"

HEADER: tuple[str, ...] = (
    "調査票番号",
    "品目番号",
    "アイテム記号",
    "調査票名",
    "項目名",
    "品目名",
    "アイテム名",
    "単位名",
)
ENGLISH_COLUMNS = 5  # Questionnaires, category, commodity, item, unit

# 調査票番号 → 産業の短い名前
SURVEYS: dict[str, str] = {"2430": "bicycle", "2400": "motor_vehicle"}

# 正規化した品目名 → product キー（design/themes/bicycle-industry.md の product 列）
PRODUCTS: dict[str, str] = {
    "軽快車": "bicycle_light",
    "電動アシスト車": "bicycle_e_assist",
    "その他": "bicycle_other",  # 2430 内の「その他」= スポーツ車等を含む完成車
    "手動式車いす": "wheelchair_manual",
    "電動式車いす": "wheelchair_electric",
    "軽自動車気筒容積660ml以下": "car_kei",
    "小型自動車気筒容積660mlを超え2000ml以下": "car_small",
    "普通自動車気筒容積2000mlを超えるもの": "car_standard",
    "二輪自動車気筒容積50ml以下": "moto_le50",
    "二輪自動車気筒容積50mlを超え125ml以下": "moto_51_125",
    "二輪自動車気筒容積125mlを超え250ml以下": "moto_126_250",
    "二輪自動車気筒容積250mlを超えるもの": "moto_gt250",
}
# 「その他」は 2430 の製品（自転車）だけに付ける（2400 にも「その他」が現れうる）
PRODUCT_SURVEY: dict[str, str] = {"その他": "2430"}
# 各ファイルに必ずあるべき product（欠けていれば様式が変わった＝ fail-closed）
REQUIRED_PRODUCTS: frozenset[str] = frozenset(
    {"bicycle_light", "bicycle_e_assist", "bicycle_other", "car_kei", "moto_gt250"}
)

MEASURES: dict[str, str] = {
    "生産数量": "production_qty",
    "生産金額": "production_value",
    "受入数量": "receipts_qty",
    "販売数量": "sales_qty",
    "販売金額": "sales_value",
    "出荷その他数量": "other_shipments_qty",
    "月末在庫数量": "inventory_qty",
    "月末従事者数": "persons_engaged",
}

_PUNCT = re.compile(r"[\s（）()・、，,.．／/]")


def normalize(name: str) -> str:
    """全角→半角（NFKC）、括弧・中点・空白を除く。年で変わる表記差を吸収する。"""
    return _PUNCT.sub("", unicodedata.normalize("NFKC", name))


def annual_url(year: int) -> str:
    return f"{BASE}/{year}/h2daa{year}k.xlsx"


def raw_name(year: int) -> str:
    return f"h2daa{year}k.xlsx"


@dataclass(frozen=True)
class Period:
    period_type: str  # month / cy / fy / quarter
    period: str  # 2025-01 / 2025 / FY2025 / 2025Q1
    year: int


def parse_time_code(code: object) -> Period:
    """時間軸コード（10 桁の int または str）→ Period。知らない形は ValueError。"""
    s = str(code).strip()
    if not re.fullmatch(r"\d{10}", s):
        msg = f"time-axis code {code!r} is not 10 digits"
        raise ValueError(msg)
    year, kind, a, b = int(s[:4]), s[4:6], s[6:8], s[8:10]
    if kind == "00" and a == b == "00":
        return Period("cy", f"{year}", year)
    if kind == "10" and a == b == "00":
        return Period("fy", f"FY{year}", year)
    if kind == "00" and a == b and 1 <= int(a) <= 12:
        return Period("month", f"{year}-{a}", year)
    if kind == "00" and (a, b) in {("01", "03"), ("04", "06"), ("07", "09"), ("10", "12")}:
        return Period("quarter", f"{year}Q{(int(b)) // 3}", year)
    msg = f"time-axis code {code!r} has an unknown layout"
    raise ValueError(msg)


def parse_value(cell: object) -> tuple[float | None, str]:
    """セル → (値, flag)。flag: '' 通常 / 'X' 秘匿 / '***' 不詳。"""
    if cell is None or (isinstance(cell, str) and cell.strip() == ""):
        return None, "missing"
    if isinstance(cell, int | float):
        return float(cell), ""
    s = str(cell).strip()
    if s == "X":
        return None, "X"
    if s == "***":
        return None, "***"
    try:
        return float(s.replace(",", "")), ""
    except ValueError as e:
        msg = f"unparseable cell {cell!r}"
        raise ValueError(msg) from e


@dataclass
class Parsed:
    rows: list[dict[str, object]]
    skipped: list[str] = field(default_factory=list)


def parse(payload: bytes, report_year: int) -> Parsed:
    """年報 xlsx の 2430・2400 行を long 形式に。様式が違えば ValueError（fail-closed）。

    行: survey_no, industry, item_no, item_code, product_ja, product_en, product, measure_ja,
    measure, unit, period_type, period, year, value, value_flag, report_year, source。
    `product` / `measure` は未知なら None（staged には残し、mart で落とす）。
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises several unrelated types for non-xlsx bytes
        msg = "payload is not an xlsx workbook"
        raise ValueError(msg) from e
    if SHEET not in wb.sheetnames:
        msg = f"sheet {SHEET!r} missing; sheets = {wb.sheetnames}"
        raise ValueError(msg)
    rows = wb[SHEET].iter_rows(values_only=True)
    try:
        code_row = next(rows)
        header_row = next(rows)
    except StopIteration as e:
        msg = "sheet has fewer than two rows"
        raise ValueError(msg) from e
    if tuple(str(c) for c in header_row[: len(HEADER)]) != HEADER:
        msg = f"header row mismatch: {header_row[: len(HEADER)]!r}"
        raise ValueError(msg)
    n_cols = len(header_row)
    first_value = len(HEADER)
    last_value = n_cols - ENGLISH_COLUMNS
    periods: list[tuple[int, Period]] = []
    for j in range(first_value, last_value):
        code = code_row[j] if j < len(code_row) else None
        if code is None:
            msg = f"time-axis code missing at column {j + 1}"
            raise ValueError(msg)
        periods.append((j, parse_time_code(code)))
    if not any(p.period_type == "month" for _, p in periods):
        msg = "no monthly columns found"
        raise ValueError(msg)

    out: list[dict[str, object]] = []
    skipped: list[str] = []
    seen_products: set[str] = set()
    unknown: set[tuple[str, str]] = set()
    for row in rows:
        survey_no = str(row[0]) if row[0] is not None else ""
        if survey_no not in SURVEYS:
            continue
        item_no, item_code = str(row[1]), str(row[2])
        product_ja, measure_ja, unit = str(row[5]), str(row[6]), str(row[7])
        key = normalize(product_ja)
        product = PRODUCTS.get(key)
        if product is not None and PRODUCT_SURVEY.get(key, survey_no) != survey_no:
            product = None
        if product is not None:
            seen_products.add(product)
        else:
            unknown.add((survey_no, product_ja))
        measure = MEASURES.get(measure_ja)
        english = row[last_value : last_value + ENGLISH_COLUMNS]
        product_en = str(english[2]) if len(english) > 2 and english[2] is not None else ""
        for j, p in periods:
            value, flag = parse_value(row[j] if j < len(row) else None)
            out.append(
                {
                    "survey_no": survey_no,
                    "industry": SURVEYS[survey_no],
                    "item_no": item_no,
                    "item_code": item_code,
                    "product_ja": product_ja,
                    "product_en": product_en,
                    "product": product,
                    "measure_ja": measure_ja,
                    "measure": measure,
                    "unit": unit,
                    "period_type": p.period_type,
                    "period": p.period,
                    "year": p.year,
                    "value": value,
                    "value_flag": flag,
                    "report_year": report_year,
                    "source": SOURCE,
                }
            )
    missing = REQUIRED_PRODUCTS - seen_products
    if missing:
        msg = f"required products missing in {report_year} workbook: {sorted(missing)}"
        raise ValueError(msg)
    skipped.extend(f"{s}: unmapped product {n!r}" for s, n in sorted(unknown))
    return Parsed(rows=out, skipped=skipped)


MART_KEY: tuple[str, ...] = ("industry", "product", "measure", "period_type", "period")


def build_mart(staged: list[dict[str, object]]) -> list[dict[str, object]]:
    """mapped な product×measure だけを残し、同じ系列×期間は最新の report_year を採る。

    月次 12 本が全部そろう年には `period_type = "cy_from_months"` の年計を足す（在庫・従事者数は
    ストックなので足さない）。出力はキー順にソート、同じキーが残れば ValueError。
    """
    best: dict[tuple[object, ...], dict[str, object]] = {}
    for r in staged:
        if r.get("product") is None or r.get("measure") is None:
            continue
        k = tuple(r[c] for c in MART_KEY)
        cur = best.get(k)
        if cur is None or int(str(r["report_year"])) > int(str(cur["report_year"])):
            best[k] = r
    out: list[dict[str, object]] = [
        {
            "iso3": "JPN",
            "industry": r["industry"],
            "product": r["product"],
            "measure": r["measure"],
            "unit": r["unit"],
            "period_type": r["period_type"],
            "period": r["period"],
            "year": r["year"],
            "value": r["value"],
            "value_flag": r["value_flag"],
            "report_year": r["report_year"],
            "source": r["source"],
        }
        for r in best.values()
    ]
    # 年計（フロー量のみ）
    flows = {"production_qty", "production_value", "receipts_qty", "sales_qty", "sales_value"}
    flows.add("other_shipments_qty")
    months: dict[tuple[object, ...], list[dict[str, object]]] = {}
    for r in out:
        if r["period_type"] == "month" and r["measure"] in flows:
            months.setdefault((r["industry"], r["product"], r["measure"], r["year"]), []).append(r)
    for (industry, product, measure, year), rs in months.items():
        if len(rs) != 12:
            continue
        values = [r["value"] for r in rs]
        complete = all(isinstance(v, float) for v in values)
        total = sum(float(v) for v in values if isinstance(v, float)) if complete else None
        out.append(
            {
                "iso3": "JPN",
                "industry": industry,
                "product": product,
                "measure": measure,
                "unit": rs[0]["unit"],
                "period_type": "cy_from_months",
                "period": f"{year}",
                "year": year,
                "value": total,
                "value_flag": "" if complete else "partial",
                "report_year": max(int(str(r["report_year"])) for r in rs),
                "source": SOURCE,
            }
        )
    out.sort(key=lambda r: tuple(str(r[c]) for c in MART_KEY))
    seen: set[tuple[object, ...]] = set()
    for r in out:
        k = tuple(r[c] for c in MART_KEY)
        if k in seen:
            msg = f"duplicate mart key {k}"
            raise ValueError(msg)
        seen.add(k)
    return out
