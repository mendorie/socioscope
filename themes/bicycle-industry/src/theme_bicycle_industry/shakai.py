"""総務省 社会生活基本調査 時系列統計表（男女・年齢・スポーツ種類別 行動者率の推移）→ long 行。

e-Stat `file-download?statInfId=000032229293&fileKind=0`（2021 年調査の公表分、appId 不要、62 KB、
シート `jikei2`）。
行 6（1-based）にスポーツ名（8 列ごと: 1986・1991・1996・2001・2006・2011・2016・2021）、
行 9 に西暦、
行 11 以降が 総数／男／女 × 年齢階級（総数、10–14 … 70 歳以上、(再掲) 15 歳以上・65 歳以上）。
値は行動者率（%）、`...` は集計なし（サイクリングは 2001 年全体、1986・1991 の総数・10–14 歳）。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import openpyxl

SOURCE = "estat_shakai"
LICENSE = (
    "政府標準利用規約（第2.0版）＝CC BY 4.0 互換。出典: 政府統計の総合窓口(e-Stat)、"
    "社会生活基本調査（総務省）令和3年 時系列統計表 を加工（確認 2026-10-22）"
)
STAT_INF_ID = "000032229293"
URL = f"https://www.e-stat.go.jp/stat-search/file-download?statInfId={STAT_INF_ID}&fileKind=0"
RAW_NAME = "shakai_seikatsu_sports_timeseries_2021.xlsx"
SHEET = "jikei2"
SEXES: dict[str, str] = {"総数": "total", "男": "male", "女": "female"}
AGES: dict[str, str] = {
    "総数": "total",
    "10～14歳": "10_14",
    "15～19歳": "15_19",
    "20～24歳": "20_24",
    "25～29歳": "25_29",
    "30～39歳": "30_39",
    "40～49歳": "40_49",
    "50～59歳": "50_59",
    "60～64歳": "60_64",
    "65～69歳": "65_69",
    "70歳以上": "70_plus",
    "(再掲)15歳以上": "r_15_plus",
    "(再掲)65歳以上": "r_65_plus",
}
REQUIRED_SPORTS: frozenset[str] = frozenset({"総数", "サイクリング"})
_WS = re.compile(r"\s+")


def _norm(s: object) -> str:
    return _WS.sub("", str(s)) if s is not None else ""


@dataclass
class Parsed:
    rows: list[dict[str, object]]
    skipped: list[str] = field(default_factory=list)


def parse(payload: bytes) -> Parsed:
    """xlsx → 行 (sport, sex, age, year, rate_pct, source)。`...` は行なし。

    様式が違えば ValueError。"""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as e:
        msg = "payload is not an xlsx workbook"
        raise ValueError(msg) from e
    if SHEET not in wb.sheetnames:
        msg = f"sheet {SHEET!r} missing; sheets = {wb.sheetnames}"
        raise ValueError(msg)
    rows = list(wb[SHEET].iter_rows(values_only=True))
    if len(rows) < 12 or _norm(rows[4][1]) != "男女年齢" or "行動者率" not in _norm(rows[4][3]):
        msg = "title cells (男女/年齢, 行動者率) not where expected"
        raise ValueError(msg)
    sport_row, year_row = rows[5], rows[8]
    # sport name → list of (column, year)
    sports: dict[str, list[tuple[int, int]]] = {}
    current: str | None = None
    for j in range(3, len(sport_row)):
        name = _norm(sport_row[j])
        if name:
            current = name
            sports[current] = []
        y = year_row[j] if j < len(year_row) else None
        if current is not None and isinstance(y, int | float):
            sports[current].append((j, int(y)))
    missing = REQUIRED_SPORTS - set(sports)
    if missing:
        msg = f"sports missing from header: {sorted(missing)}"
        raise ValueError(msg)
    out: list[dict[str, object]] = []
    skipped: list[str] = []
    sex: str | None = None
    seen_cells: set[tuple[str, str]] = set()
    for r in rows[10:]:
        label = _norm(r[1]) if len(r) > 1 else ""
        if not label:
            continue
        if label in SEXES:
            sex = SEXES[label]
            age = "total"
        elif label in AGES and sex is not None:
            age = AGES[label]
        else:
            skipped.append(f"unrecognised row label {label!r}")
            continue
        seen_cells.add((sex, age))
        for sport, cols in sports.items():
            for j, year in cols:
                v = r[j] if j < len(r) else None
                if v is None or (isinstance(v, str) and v.strip() in {"...", "-", ""}):
                    continue
                if isinstance(v, str):
                    msg = f"non-numeric cell {v!r} at {sport}/{sex}/{age}/{year}"
                    raise ValueError(msg)
                out.append(
                    {
                        "sport": sport,
                        "sex": sex,
                        "age": age,
                        "year": year,
                        "rate_pct": float(v),
                        "source": SOURCE,
                    }
                )
    expected = {(s, a) for s in SEXES.values() for a in AGES.values()}
    if not expected <= seen_cells:
        msg = f"sex×age rows missing: {sorted(expected - seen_cells)[:5]}"
        raise ValueError(msg)
    return Parsed(rows=out, skipped=skipped)


MART_KEY: tuple[str, ...] = ("sport", "sex", "age", "year")


def build_mart(staged: list[dict[str, object]]) -> list[dict[str, object]]:
    out = [{"iso3": "JPN", **r} for r in staged]
    out.sort(key=lambda r: tuple(str(r[c]) for c in MART_KEY))
    seen: set[tuple[object, ...]] = set()
    for r in out:
        k = tuple(r[c] for c in MART_KEY)
        if k in seen:
            msg = f"duplicate shakai key {k}"
            raise ValueError(msg)
        seen.add(k)
    return out
