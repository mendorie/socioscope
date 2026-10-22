from pathlib import Path

import openpyxl
import pytest

from theme_bicycle_industry import census

FIXTURE = (Path(__file__).parent / "fixtures" / "census_f1_001_sample.xlsx").read_bytes()


def test_url_and_names() -> None:
    assert census.URL.endswith("statInfId=000040032840&fileKind=0")
    assert census.RAW_NAME == "census_r3_retail_table1.xlsx"


def test_parse_yields_bicycle_retail_time_series_from_the_national_rows() -> None:
    parsed = census.parse(FIXTURE)
    rows = parsed.rows
    assert {r["industry"] for r in rows} == set(census.INDUSTRIES.values())
    bike = {(r["year"], r["metric"]): r["value"] for r in rows if r["industry"] == "bicycle_retail"}
    assert bike[(2016, "establishments")] == 11207 and bike[(2021, "establishments")] == 9552
    assert bike[(2021, "establishments_individual")] == 6628
    assert bike[(2021, "annual_sales_mn_jpy")] == 231252
    assert bike[(1994, "establishments")] == 17724
    assert sorted({r["year"] for r in rows}) == sorted(census.SURVEY_OF_YEAR)
    assert {r["survey"] for r in rows if r["year"] == 2014} == {"commercial_census"}
    # the 東京都 rows in the fixture are not national and are ignored
    assert all(r["industry_label"] in census.INDUSTRIES for r in rows)
    assert len([r for r in rows if r["industry"] == "bicycle_retail"]) == 10 * len(census.METRICS)
    assert parsed.skipped == []


def test_cell_handles_suppressed_and_missing_markers() -> None:
    assert census._cell("-") is None and census._cell("x") is None and census._cell(None) is None
    assert census._cell("1,234") == 1234.0 and census._cell(5) == 5.0
    with pytest.raises(ValueError, match="non-numeric"):
        census._cell("n/a")


def test_parse_fails_closed_on_html_wrong_sheet_or_missing_industry(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="not an xlsx"):
        census.parse(b"<html>404</html>")
    wb = openpyxl.load_workbook(Path(__file__).parent / "fixtures" / "census_f1_001_sample.xlsx")
    wb.active.title = "other"
    p = tmp_path / "x.xlsx"
    wb.save(p)
    with pytest.raises(ValueError, match="sheet 'f1_001' missing"):
        census.parse(p.read_bytes())
    wb = openpyxl.load_workbook(Path(__file__).parent / "fixtures" / "census_f1_001_sample.xlsx")
    ws = wb.active
    for row in ws.iter_rows():
        if row[2].value == "592_自転車小売業":
            row[2].value = "592_?"
    wb.save(p)
    with pytest.raises(ValueError, match="industries missing"):
        census.parse(p.read_bytes())


def test_build_mart_sorts_and_rejects_duplicates() -> None:
    rows = census.parse(FIXTURE).rows
    mart = census.build_mart(rows)
    assert all(r["iso3"] == "JPN" for r in mart)
    keys = [tuple(r[c] for c in census.MART_KEY) for r in mart]
    assert keys == sorted(keys) and len(set(keys)) == len(keys)
    with pytest.raises(ValueError, match="duplicate census key"):
        census.build_mart([*rows, dict(rows[0])])
