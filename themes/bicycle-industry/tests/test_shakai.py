from pathlib import Path

import openpyxl
import pytest

from theme_bicycle_industry import shakai

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "shakai_jikei2_sample.xlsx"
FIXTURE = FIXTURE_PATH.read_bytes()


def test_url_and_constants() -> None:
    assert shakai.URL.endswith("statInfId=000032229293&fileKind=0")
    assert shakai.RAW_NAME.endswith(".xlsx") and shakai.SHEET == "jikei2"


def test_parse_yields_sport_sex_age_year_rates_and_drops_dots() -> None:
    parsed = shakai.parse(FIXTURE)
    rows = parsed.rows
    assert {r["sport"] for r in rows} >= {"総数", "サイクリング", "ジョギング・マラソン"}
    cyc = {
        (r["sex"], r["age"], r["year"]): r["rate_pct"] for r in rows if r["sport"] == "サイクリング"
    }
    assert cyc[("total", "total", 2021)] == 8.2 and cyc[("total", "total", 1996)] == 7.6
    assert cyc[("male", "20_24", 2021)] == 14.1 and cyc[("female", "70_plus", 2021)] == 1.6
    assert ("total", "total", 2001) not in cyc and ("total", "total", 1986) not in cyc  # '...'
    assert ("total", "15_19", 1986) in cyc
    years = {r["year"] for r in rows if r["sport"] == "サイクリング"}
    assert years == {1986, 1991, 1996, 2006, 2011, 2016, 2021}
    assert {r["age"] for r in rows} == set(shakai.AGES.values())
    assert parsed.skipped == []


def test_parse_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="not an xlsx"):
        shakai.parse(b"<html>")
    wb = openpyxl.load_workbook(FIXTURE_PATH)
    ws = wb.active
    for row in ws.iter_rows(min_row=6, max_row=6):
        for c in row:
            if c.value == "サイクリング":
                c.value = "自転車"
    p = tmp_path / "x.xlsx"
    wb.save(p)
    with pytest.raises(ValueError, match="sports missing"):
        shakai.parse(p.read_bytes())
    wb = openpyxl.load_workbook(FIXTURE_PATH)
    wb.active.title = "other"
    wb.save(p)
    with pytest.raises(ValueError, match="sheet 'jikei2' missing"):
        shakai.parse(p.read_bytes())


def test_build_mart_sorts_and_rejects_duplicates() -> None:
    rows = shakai.parse(FIXTURE).rows
    mart = shakai.build_mart(rows)
    keys = [tuple(r[c] for c in shakai.MART_KEY) for r in mart]
    assert keys == sorted(keys) and len(set(keys)) == len(keys)
    assert all(r["iso3"] == "JPN" for r in mart)
    with pytest.raises(ValueError, match="duplicate shakai key"):
        shakai.build_mart([*rows, dict(rows[0])])
