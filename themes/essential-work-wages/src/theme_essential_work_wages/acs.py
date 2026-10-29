"""US Census ACS 1-year: median earnings and workers by detailed occupation (B24121-B24126).

Source: table-based summary files of the ACS 1-year estimates (public domain, no API key).
  data   https://www2.census.gov/programs-surveys/acs/summary_file/<YEAR>/table-based-SF/data/1YRData/acsdt1y<YEAR>-b2412N.dat
  shells https://www2.census.gov/programs-surveys/acs/summary_file/2023/table-based-SF/documentation/ACS20231YR_Table_Shells.txt
  codes  https://www2.census.gov/programs-surveys/demo/guidance/industry-occupation/2018-occupation-code-list-and-crosswalk.xlsx

Tables (universe: full-time, year-round civilian employed population 16+ with earnings):
  B24121 median earnings, total   B24122 male   B24123 female
  B24124 workers, total           B24125 male   B24126 female
Each has 566 lines: 001 = Total, 002-566 = the 565 detailed civilian occupations of the 2018
Census occupation code list (military 98xx and 9920 are excluded). A ``.dat`` file is
pipe-separated: header ``GEO_ID|B2412N_E001|B2412N_M001|...`` and a single national row
``0100000US``. ``E`` = estimate, ``M`` = margin of error (90 %). Negative values are Census
sentinels (-666666666 not available, -333333333 median in an open-ended interval,
-222222222 MOE not computable, ...) and become ``None``.

Years: 2021-2024 (the table-based format starts in 2021; the 2020 1-year release was not
published because of the pandemic). The ``Total`` line is kept as ``occupation_code='0000'``.

Labels: the B24121 line labels are identical in the 2021, 2022, 2023 and 2024 shells
(checked field by field), so one shells file (2023) is fetched. Labels are matched to the code
table by title (lower case, punctuation collapsed) rather than by position, because the shells
order and the code-list order differ in one spot (8865 "Other production equipment operators
and tenders" sits after 8950 in the shells). With that normalisation all 565 labels match the
2018 Census titles exactly; any miss raises ``ValueError`` listing the labels.

Code table quirks handled here: the "Summary of 2018 Changes" sheet repeats code 0650 twice
(identical rows); the "2010 to 2018 Crosswalk " sheet is block-structured (a 2010 row with
blank 2018 columns followed by its 2018 successors; note rows carry only a title in column F)
and lacks the brand-new code 8225 (51-4XXX). The 2010 SOC codes are kept as
``soc_2010_codes`` for sources still on SOC 2010 (Dingel-Neiman).

Caveats for analysis: medians are nominal dollars of each survey year (the ACS "inflation
adjusted" note refers to within-year adjustment only); only full-time year-round workers are
covered, so part-time heavy occupations are under-represented; 1-year estimates for small
occupations have large MOEs (``earnings_median_moe``).
"""

import io
import re
from collections import defaultdict

import openpyxl

from socioscope_core.core.pipeline import Context
from socioscope_core.ports.fetcher import FetchError
from theme_essential_work_wages.pipeline import THEME

SOURCE = "census_acs1"
LICENSE = "Public domain (U.S. Census Bureau, American Community Survey 1-year summary file)"
CLASSIFICATION = "us_census_occ2018"

YEARS: tuple[int, ...] = (2021, 2022, 2023, 2024)
TABLES: tuple[str, ...] = ("B24121", "B24122", "B24123", "B24124", "B24125", "B24126")
SHELLS_YEAR = 2023

_BASE = "https://www2.census.gov/programs-surveys/acs/summary_file"
SHELLS_RAW_NAME = f"ACS{SHELLS_YEAR}1YR_Table_Shells.txt"
SHELLS_URL = f"{_BASE}/{SHELLS_YEAR}/table-based-SF/documentation/{SHELLS_RAW_NAME}"
CODES_RAW_NAME = "2018-occupation-code-list-and-crosswalk.xlsx"
CODES_URL = (
    "https://www2.census.gov/programs-surveys/demo/guidance/industry-occupation/" + CODES_RAW_NAME
)

MAGIC = b"GEO_ID|"
SHELLS_MAGIC = b"Table ID|"
XLSX_MAGIC = b"PK\x03\x04"
_BOM = b"\xef\xbb\xbf"

CODES_TABLE = "staged/acs/occupation_codes"
MART_TABLE = "marts/acs_occupation_earnings"
NATIONAL_GEO_ID = "0100000US"
TOTAL_CODE = "0000"
_MILITARY_FROM = "9800"
_SOC = re.compile(r"\d{2}-[\dX]{4}|none")
_SUMMARY_SHEET = "Summary of 2018 Changes"
_XWALK_SHEET = "2010 to 2018 Crosswalk "

# mart column <- (table, E or M)
_MART_VALUES: tuple[tuple[str, str], ...] = (
    ("earnings_median", "b24121_e"),
    ("earnings_median_male", "b24122_e"),
    ("earnings_median_female", "b24123_e"),
    ("workers", "b24124_e"),
    ("workers_male", "b24125_e"),
    ("workers_female", "b24126_e"),
    ("earnings_median_moe", "b24121_m"),
)


def url(year: int, table: str) -> str:
    return f"{_BASE}/{year}/table-based-SF/data/1YRData/{raw_name(year, table)}"


def raw_name(year: int, table: str) -> str:
    return f"acsdt1y{year}-{table.lower()}.dat"


def staged_table(year: int) -> str:
    return f"staged/acs/occupation_earnings_{year}"


def is_civilian(code: str) -> bool:
    return code < _MILITARY_FROM


def normalize_title(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


# ---- fetch


def _fetch_one(
    ctx: Context, name: str, u: str, magic: bytes, kind: str, written: list[str], skipped: list[str]
) -> None:
    try:
        payload = ctx.fetcher.fetch(u)
    except FetchError as e:
        skipped.append(f"{name}: {e}")
        return
    if not payload.removeprefix(_BOM).startswith(magic):
        skipped.append(f"{name}: non-{kind} payload")
        return
    rec = ctx.raw.put(
        theme=THEME, source=SOURCE, name=name, url=u, license=LICENSE, payload=payload
    )
    written.append(rec.relative_path)


def fetch(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    for year in YEARS:
        for table in TABLES:
            _fetch_one(ctx, raw_name(year, table), url(year, table), MAGIC, "dat", written, skipped)
    _fetch_one(ctx, SHELLS_RAW_NAME, SHELLS_URL, SHELLS_MAGIC, "shells", written, skipped)
    _fetch_one(ctx, CODES_RAW_NAME, CODES_URL, XLSX_MAGIC, "xlsx", written, skipped)
    return written, skipped


# ---- pure parsers


def _value(cell: str) -> float | None:
    cell = cell.strip()
    if cell in ("", "-", "(X)", "N", "null"):
        return None
    try:
        v = float(cell)
    except ValueError as e:
        msg = f"non-numeric cell {cell!r}"
        raise ValueError(msg) from e
    return None if v < 0 else v


def parse_dat(
    payload: bytes, year: int, table: str
) -> dict[int, tuple[float | None, float | None]]:
    """Line number -> (estimate, margin of error) for the national row of one table."""
    lines = payload.decode("utf-8").splitlines()
    if len(lines) != 2:
        msg = f"{table} {year}: expected header + 1 national row, got {len(lines)} lines"
        raise ValueError(msg)
    header = lines[0].split("|")
    cells = lines[1].split("|")
    if header[0] != "GEO_ID" or len(header) != len(cells) or len(header) % 2 == 0:
        msg = f"{table} {year}: malformed header/row ({len(header)}/{len(cells)} columns)"
        raise ValueError(msg)
    if cells[0] != NATIONAL_GEO_ID:
        msg = f"{table} {year}: expected GEO_ID {NATIONAL_GEO_ID}, got {cells[0]!r}"
        raise ValueError(msg)
    out: dict[int, tuple[float | None, float | None]] = {}
    for i in range(1, len(header), 2):
        m = re.fullmatch(rf"{table}_E(\d{{3}})", header[i])
        if m is None or header[i + 1] != f"{table}_M{m.group(1)}":
            msg = f"{table} {year}: unexpected columns {header[i]!r}, {header[i + 1]!r}"
            raise ValueError(msg)
        out[int(m.group(1))] = (_value(cells[i]), _value(cells[i + 1]))
    return out


def parse_shells(payload: bytes, table: str) -> list[tuple[int, str]]:
    """(line, label) for one table, in file order."""
    text = payload.removeprefix(_BOM).decode("utf-8")
    out: list[tuple[int, str]] = []
    for raw in text.splitlines()[1:]:
        cols = raw.split("|")
        if len(cols) >= 5 and cols[0] == table:
            out.append((int(float(cols[1])), cols[4].strip()))
    if not out:
        msg = f"table {table} not found in shells"
        raise ValueError(msg)
    return out


def _sheet_rows(payload: bytes, sheet: str, width: int) -> list[tuple[str | None, ...]]:
    wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        msg = f"sheet {sheet!r} not found (have {wb.sheetnames})"
        raise ValueError(msg)
    rows: list[tuple[str | None, ...]] = []
    for r in wb[sheet].iter_rows(values_only=True):
        cells = tuple(str(c).strip() if c is not None else None for c in r[:width])
        rows.append(cells + (None,) * (width - len(cells)))
    return rows


def parse_codes(payload: bytes) -> list[tuple[str, str, str]]:
    """(census_code, soc_code, title) from the 'Summary of 2018 Changes' sheet, by census code."""
    found: dict[str, tuple[str, str, str]] = {}
    for soc, census, title, *_ in _sheet_rows(payload, _SUMMARY_SHEET, 3):
        if not (soc and census and title and _SOC.fullmatch(soc) and census.isdigit()):
            continue
        row = (census, soc, title)
        if found.setdefault(census, row) != row:
            msg = f"conflicting rows for census code {census}: {found[census]} vs {row}"
            raise ValueError(msg)
    if not found:
        msg = f"no occupation rows in sheet {_SUMMARY_SHEET!r}"
        raise ValueError(msg)
    return [found[k] for k in sorted(found)]


def parse_crosswalk(payload: bytes) -> dict[str, tuple[str, ...]]:
    """2018 census code -> sorted 2010 SOC codes, from the block-structured crosswalk sheet."""
    out: dict[str, set[str]] = defaultdict(set)
    carry: list[str] = []
    after_2018_only = False
    for s10, _c10, _t10, s18, c18, _t18 in _sheet_rows(payload, _XWALK_SHEET, 6):
        has10 = bool(s10 and _SOC.fullmatch(s10))
        has18 = bool(s18 and c18 and _SOC.fullmatch(s18) and c18.isdigit())
        if has10 and has18:
            assert s10 is not None and c18 is not None
            out[c18].add(s10)
            carry, after_2018_only = [], False
        elif has10:
            assert s10 is not None
            if after_2018_only:
                carry = []
            carry.append(s10)
            after_2018_only = False
        elif has18:
            assert c18 is not None
            out[c18].update(carry)
            after_2018_only = True
    if not out:
        msg = f"no rows in sheet {_XWALK_SHEET!r}"
        raise ValueError(msg)
    return {k: tuple(sorted(v)) for k, v in out.items()}


def build_codes_table(
    codes: list[tuple[str, str, str]], xwalk: dict[str, tuple[str, ...]]
) -> list[dict[str, object]]:
    """Staged civilian occupation code list with the 2010 SOC codes joined by ';'."""
    return [
        {
            "occupation_code": census,
            "occupation": title,
            "soc_code": soc,
            "soc_2010_codes": ";".join(xwalk.get(census, ())),
            "classification": CLASSIFICATION,
        }
        for census, soc, title in codes
        if is_civilian(census)
    ]


def build_staged(
    year: int,
    shells_lines: list[tuple[int, str]],
    parsed: dict[str, dict[int, tuple[float | None, float | None]]],
) -> list[dict[str, object]]:
    """One wide row per shells line: year, line, label, <table>_e / <table>_m for every table."""
    lines = [ln for ln, _ in shells_lines]
    if len(set(lines)) != len(lines) or set(parsed) != set(TABLES):
        msg = f"{year}: duplicate shells lines or tables {sorted(parsed)} != {TABLES}"
        raise ValueError(msg)
    for table, values in parsed.items():
        if set(values) != set(lines):
            msg = (
                f"{year} {table}: lines {sorted(values)[:3]}... do not match shells {lines[:3]}..."
            )
            raise ValueError(msg)
    out: list[dict[str, object]] = []
    for ln, label in shells_lines:
        row: dict[str, object] = {"year": year, "line": ln, "label": label}
        for table in TABLES:
            e, m = parsed[table][ln]
            row[f"{table.lower()}_e"] = e
            row[f"{table.lower()}_m"] = m
        out.append(row)
    return out


def build_mart(
    staged_rows: list[dict[str, object]], codes: list[dict[str, object]]
) -> list[dict[str, object]]:
    """Year x occupation with earnings/workers columns; labels joined to codes by title."""
    by_title: dict[str, dict[str, object]] = {}
    for c in codes:
        key = normalize_title(str(c["occupation"]))
        if key in by_title:
            msg = f"duplicate normalised title in code table: {c['occupation']!r}"
            raise ValueError(msg)
        by_title[key] = c
    years: dict[int, list[dict[str, object]]] = defaultdict(list)
    for r in staged_rows:
        years[int(str(r["year"]))].append(r)
    out: list[dict[str, object]] = []
    for year in sorted(years):
        rows = sorted(years[year], key=lambda r: int(str(r["line"])))
        lines = [int(str(r["line"])) for r in rows]
        if lines != list(range(1, len(rows) + 1)):
            msg = f"{year}: lines are not 1..n without gaps or duplicates ({lines[:5]}...)"
            raise ValueError(msg)
        if len(rows) - 1 != len(codes):
            msg = f"{year}: {len(rows) - 1} occupation lines vs {len(codes)} civilian codes"
            raise ValueError(msg)
        unmatched = [
            str(r["label"]) for r in rows[1:] if normalize_title(str(r["label"])) not in by_title
        ]
        if unmatched:
            msg = f"{year}: {len(unmatched)} shells labels not in code table: {unmatched}"
            raise ValueError(msg)
        for r in rows:
            if int(str(r["line"])) == 1:
                code, title, soc = TOTAL_CODE, "Total", None
            else:
                c = by_title[normalize_title(str(r["label"]))]
                code, title, soc = (
                    str(c["occupation_code"]),
                    str(c["occupation"]),
                    str(c["soc_code"]),
                )
            row: dict[str, object] = {
                "iso3": "USA",
                "year": year,
                "occupation_code": code,
                "occupation": title,
                "classification": CLASSIFICATION,
                "soc_code": soc,
            }
            for col, src in _MART_VALUES:
                row[col] = r[src]
            row["source"] = SOURCE
            out.append(row)
    out.sort(key=lambda r: (int(str(r["year"])), str(r["occupation_code"])))
    return out


# ---- stage / mart


def stage(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    shells_payload = ctx.raw.get(theme=THEME, source=SOURCE, name=SHELLS_RAW_NAME)
    if shells_payload is None:
        skipped.append(f"{SHELLS_RAW_NAME}: no raw")
    codes_payload = ctx.raw.get(theme=THEME, source=SOURCE, name=CODES_RAW_NAME)
    if codes_payload is None:
        skipped.append(f"{CODES_RAW_NAME}: no raw")
    else:
        ctx.tables.write_table(
            CODES_TABLE,
            build_codes_table(parse_codes(codes_payload), parse_crosswalk(codes_payload)),
        )
        written.append(CODES_TABLE)
    for year in YEARS:
        payloads: dict[str, bytes] = {}
        for table in TABLES:
            p = ctx.raw.get(theme=THEME, source=SOURCE, name=raw_name(year, table))
            if p is None:
                skipped.append(f"{raw_name(year, table)}: no raw")
            else:
                payloads[table] = p
        if shells_payload is None or len(payloads) != len(TABLES):
            continue
        parsed = {t: parse_dat(p, year, t) for t, p in payloads.items()}
        ctx.tables.write_table(
            staged_table(year), build_staged(year, parse_shells(shells_payload, "B24121"), parsed)
        )
        written.append(staged_table(year))
    return written, skipped


def mart(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    try:
        codes = ctx.tables.read_table(CODES_TABLE)
    except FileNotFoundError:
        skipped.append(f"{CODES_TABLE}: no staged table")
        return written, skipped
    rows: list[dict[str, object]] = []
    for year in YEARS:
        try:
            rows.extend(ctx.tables.read_table(staged_table(year)))
        except FileNotFoundError:
            skipped.append(f"{staged_table(year)}: no staged table")
    if rows:
        ctx.tables.write_table(MART_TABLE, build_mart(rows, codes))
        written.append(MART_TABLE)
    return written, skipped
