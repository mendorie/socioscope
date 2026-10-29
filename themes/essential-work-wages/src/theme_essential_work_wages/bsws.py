"""賃金構造基本統計調査（厚生労働省、e-Stat）一般労働者・職種別賃金 2009-2025 年。

採用した表（毎年 1 ファイル、産業計・企業規模 4 区分・性 3 区分・職種別の「（職種）第 1 表」）:

- 2020 年調査以降（新職種分類、``classification='jp_bsws2020'``、145 職種）:
  「職種（小分類）別きまって支給する現金給与額、所定内給与額及び
  年間賞与その他特別給与額（産業計）」。同ページの「【参考】…（役職者を除く）」ではなく
  本表（e-Stat 一覧の先頭、大分類表の 1 つ前の ID）。2020 年は xls（fileKind=0）、
  2021 年以降は xlsx（fileKind=4）。1 シート、男女計→男→女 の 3 ブロック。
- 2009-2019 年調査（旧職種分類、``classification='jp_bsws2009'``、129 職種）:
  「職種別第 1 表 職種別きまって支給する現金給与額、所定内給与額及び
  年間賞与その他特別給与額（産業計）」。同ページの「職種・性、年齢階級別…」は
  年齢階級別の長い表で企業規模計しか持たないので、列構成が新形式と同じ
  （企業規模 4 区分 × 8 指標）で性別 3 ブロックを持つ第 1 表を採る。すべて xls（fileKind=0）。
  2010-2019 年は 1 シート 3 ブロック、2009 年は 男女計／男／女 の 3 シート
  （男女計シートは独立した性別ラベル行を持つ）。

2020 年の職種分類改定で旧 129 職種と新 145 職種は接続しない。ここでは接続せず
``classification`` 列で区別する（対応表経由の接続はテーマ側の設計に従う）。

表の形: 各行 = 職種、列 = 企業規模計(10 人以上)／1,000 人以上／100〜999 人／10〜99 人 ×
(年齢[歳], 勤続年数[年], 所定内実労働時間数[時間], 超過実労働時間数[時間],
きまって支給する現金給与額[千円], 所定内給与額[千円], 年間賞与その他特別給与額[千円],
労働者数[十人])。労働者数は 10 倍して人数にする。先頭セルには
"　男女計\\n\\n管理的職業従事者" のように性別ラベルが混ざる（ブロック先頭と、
印刷ページ送りの位置で再掲）。'-' は該当労働者なし／公表なしで None。

ライセンス: 政府標準利用規約（第 2.0 版）＝ CC BY 4.0 互換。
出典表示: 厚生労働省「賃金構造基本統計調査」。
"""

import io
import re
from collections.abc import Sequence

from socioscope_core.core.pipeline import Context
from socioscope_core.ports.fetcher import FetchError
from theme_essential_work_wages.pipeline import THEME

SOURCE = "estat_bsws"
LICENSE = "政府標準利用規約（第2.0版）（CC BY 4.0 互換）: 厚生労働省「賃金構造基本統計調査」"
DOWNLOAD_URL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId={stat_inf_id}&fileKind={file_kind}"

# year -> (statInfId, fileKind). fileKind 4 = xlsx, 0 = xls. See module docstring for the choice.
TABLES: dict[int, tuple[str, int]] = {
    2025: ("000040421116", 4),
    2024: ("000040247854", 4),
    2023: ("000040163741", 4),
    2022: ("000040029181", 4),
    2021: ("000032182952", 4),
    2020: ("000032069431", 0),
    2019: ("000031919843", 0),
    2018: ("000031683128", 0),
    2017: ("000031559736", 0),
    2016: ("000031227940", 0),
    2015: ("000028996524", 0),
    2014: ("000026308728", 0),
    2013: ("000023602572", 0),
    2012: ("000016347004", 0),
    2011: ("000012743013", 0),
    2010: ("000008278745", 0),
    2009: ("000007214617", 0),
}

MAGIC_XLSX = b"PK\x03\x04"
MAGIC_XLS = b"\xd0\xcf\x11\xe0"
MAGIC: tuple[bytes, ...] = (MAGIC_XLSX, MAGIC_XLS)

STAGED_TABLE = "staged/bsws/occupation_wages_{year}"
MART_TABLE = "marts/bsws_occupation_wages"

EXPECTED_OCCUPATIONS: dict[str, int] = {"jp_bsws2020": 145, "jp_bsws2009": 129}

SIZE_CLASSES: tuple[tuple[str, str], ...] = (
    ("企業規模計（10人以上）", "all"),
    ("1,000人以上", "1000+"),
    ("100～999人", "100-999"),
    ("10～99人", "10-99"),
)
METRIC_COLUMNS: tuple[str, ...] = (
    "age",
    "tenure_years",
    "scheduled_hours",
    "overtime_hours",
    "monthly_cash_thousand_yen",
    "scheduled_wage_thousand_yen",
    "annual_bonus_thousand_yen",
    "workers",
)
UNITS: tuple[str, ...] = ("歳", "年", "時間", "時間", "千円", "千円", "千円", "十人")
STAGED_COLUMNS: tuple[str, ...] = (
    "year",
    "classification",
    "sex",
    "size_class",
    "occupation",
    *METRIC_COLUMNS,
)
MART_COLUMNS: tuple[str, ...] = ("iso3", *STAGED_COLUMNS, "source")
MART_KEY: tuple[str, ...] = ("year", "sex", "size_class", "occupation")

SEXES: tuple[tuple[str, str], ...] = (("男女計", "T"), ("男", "M"), ("女", "F"))
_SEX_PREFIX = re.compile(r"^[\s　]*(男女計|男|女)[\s　]*(?:\n+|$)")
_MISSING = {"-", "X", "x", "…", "－", ""}

Grid = list[list[object]]


def file_kind_ext(file_kind: int) -> str:
    return "xlsx" if file_kind == 4 else "xls"


def raw_name(year: int) -> str:
    return f"bsws_{year}_occupation.{file_kind_ext(TABLES[year][1])}"


def url(year: int) -> str:
    stat_inf_id, file_kind = TABLES[year]
    return DOWNLOAD_URL.format(stat_inf_id=stat_inf_id, file_kind=file_kind)


def classification(year: int) -> str:
    return "jp_bsws2020" if year >= 2020 else "jp_bsws2009"


def staged_table(year: int) -> str:
    return STAGED_TABLE.format(year=year)


# ---- fetch / stage / mart ---------------------------------------------------------


def fetch(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    for year in sorted(TABLES):
        name = raw_name(year)
        try:
            payload = ctx.fetcher.fetch(url(year))
        except FetchError as e:
            skipped.append(f"{name}: {e}")
            continue
        if not payload.startswith(MAGIC):
            skipped.append(f"{name}: non-excel payload")
            continue
        rec = ctx.raw.put(
            theme=THEME, source=SOURCE, name=name, url=url(year), license=LICENSE, payload=payload
        )
        written.append(rec.relative_path)
    return written, skipped


def stage(
    ctx: Context,
    written: list[str],
    skipped: list[str],
    *,
    expected_occupations: int | None = None,
) -> tuple[list[str], list[str]]:
    for year in sorted(TABLES):
        name = raw_name(year)
        payload = ctx.raw.get(theme=THEME, source=SOURCE, name=name)
        if payload is None:
            skipped.append(f"{name}: no raw")
            continue
        rows = parse(payload, year, expected_occupations=expected_occupations)
        table = staged_table(year)
        ctx.tables.write_table(table, rows)
        written.append(table)
    return written, skipped


def mart(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    rows: list[dict[str, object]] = []
    for year in sorted(TABLES):
        try:
            rows.extend(ctx.tables.read_table(staged_table(year)))
        except FileNotFoundError:
            continue
    if not rows:
        skipped.append(f"{MART_TABLE}: no staged tables")
        return written, skipped
    ctx.tables.write_table(MART_TABLE, build_mart(rows))
    written.append(MART_TABLE)
    return written, skipped


# ---- pure transforms ---------------------------------------------------------------


def parse(
    payload: bytes, year: int, *, expected_occupations: int | None = None
) -> list[dict[str, object]]:
    """Excel bytes -> wide rows (year x sex x size_class x occupation). Pure given the bytes."""
    return parse_grids(_load_grids(payload), year, expected_occupations=expected_occupations)


def parse_grids(
    grids: Sequence[Grid], year: int, *, expected_occupations: int | None = None
) -> list[dict[str, object]]:
    """Cell grids (one per sheet) -> wide rows. Raises ValueError when the layout drifts."""
    cls = classification(year)
    expected = EXPECTED_OCCUPATIONS[cls] if expected_occupations is None else expected_occupations
    entries: list[tuple[str | None, str, list[object]]] = []
    for grid in grids:
        entries.extend(_sheet_entries(grid))
    if len(entries) != 3 * expected:
        msg = f"{year}: expected {expected} occupations x 3 sexes, found {len(entries)} data rows"
        raise ValueError(msg)
    blocks = [entries[i * expected : (i + 1) * expected] for i in range(3)]
    names = [e[1] for e in blocks[0]]
    if len(set(names)) != len(names):
        msg = f"{year}: duplicate occupation names in the first block"
        raise ValueError(msg)
    out: list[dict[str, object]] = []
    for (sex_label, sex), block in zip(SEXES, blocks, strict=True):
        if [e[1] for e in block] != names:
            msg = f"{year}: occupation list of block {sex!r} differs from the first block"
            raise ValueError(msg)
        for prefix, occupation, values in block:
            if prefix is not None and prefix != sex_label:
                msg = f"{year}: sex prefix {prefix!r} inside block {sex_label!r} ({occupation})"
                raise ValueError(msg)
            for i, (_, size_class) in enumerate(SIZE_CLASSES):
                cells = values[i * len(UNITS) : (i + 1) * len(UNITS)]
                row: dict[str, object] = {
                    "year": year,
                    "classification": cls,
                    "sex": sex,
                    "size_class": size_class,
                    "occupation": occupation,
                }
                for col, unit, cell in zip(METRIC_COLUMNS, UNITS, cells, strict=True):
                    row[col] = _number(cell, unit, f"{year} {occupation} {size_class} {col}")
                out.append(row)
    return out


def build_mart(rows: Sequence[dict[str, object]]) -> list[dict[str, object]]:
    """Add iso3/source, keep the mart columns, sort by key and reject duplicate keys."""
    out = [{c: {"iso3": "JPN", "source": SOURCE, **r}.get(c) for c in MART_COLUMNS} for r in rows]
    out.sort(key=lambda r: tuple(str(r[k]) for k in MART_KEY))
    seen: set[tuple[object, ...]] = set()
    for r in out:
        k = tuple(r[c] for c in MART_KEY)
        if k in seen:
            msg = f"duplicate key {k}"
            raise ValueError(msg)
        seen.add(k)
    return out


# ---- helpers -------------------------------------------------------------------------


def _load_grids(payload: bytes) -> list[Grid]:
    if payload.startswith(MAGIC_XLSX):
        import openpyxl

        try:
            wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        except Exception as e:
            msg = "payload is not a readable xlsx"
            raise ValueError(msg) from e
        return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    if payload.startswith(MAGIC_XLS):
        import xlrd

        try:
            book = xlrd.open_workbook(file_contents=payload)
        except Exception as e:
            msg = "payload is not a readable xls"
            raise ValueError(msg) from e
        return [[sh.row_values(r) for r in range(sh.nrows)] for sh in book.sheets()]
    msg = "payload is neither xlsx nor xls"
    raise ValueError(msg)


def _text(cell: object) -> str:
    return cell.strip() if isinstance(cell, str) else ""


def _sheet_entries(grid: Grid) -> list[tuple[str | None, str, list[object]]]:
    """One sheet -> [(sex_prefix | None, occupation, 32 metric cells)] in sheet order.

    Locates the header by its labels (区分 / 企業規模 / units) instead of fixed positions, so the
    one-row offset between years and the leading code columns of the old xls do not matter.
    """
    unit_row = next((i for i, r in enumerate(grid) if "十人" in r), None)
    if unit_row is None or unit_row < 3:
        msg = "unit row (歳/年/時間/千円/十人) not found"
        raise ValueError(msg)
    value_cols = [j for j, c in enumerate(grid[unit_row]) if _text(c)]
    units = [_text(grid[unit_row][j]) for j in value_cols]
    if units != list(UNITS) * len(SIZE_CLASSES):
        msg = f"unexpected unit row {units}"
        raise ValueError(msg)
    size_row = grid[unit_row - 3]
    sizes = [_text(c) for c in size_row if _text(c)]
    label_col = next((j for j, c in enumerate(size_row) if _text(c).startswith("区")), None)
    if label_col is None or sizes[1:] != [s for s, _ in SIZE_CLASSES]:
        msg = f"unexpected size-class header {sizes}"
        raise ValueError(msg)
    entries: list[tuple[str | None, str, list[object]]] = []
    for r in grid[unit_row + 1 :]:
        raw_label = r[label_col] if label_col < len(r) else None
        if not isinstance(raw_label, str) or not raw_label.strip():
            continue
        m = _SEX_PREFIX.match(raw_label)
        prefix = m.group(1) if m else None
        occupation = raw_label[m.end() :] if m else raw_label
        occupation = occupation.replace("　", "").replace("\n", "").strip()
        if not occupation:
            continue  # standalone sex label row (2009 男女計 sheet)
        values = [r[j] if j < len(r) else None for j in value_cols]
        entries.append((prefix, occupation, values))
    return entries


def _number(cell: object, unit: str, where: str) -> float | int | None:
    if cell is None or (isinstance(cell, str) and cell.strip() in _MISSING):
        return None
    if isinstance(cell, bool) or not isinstance(cell, (int, float)):
        msg = f"unexpected cell value {cell!r} at {where}"
        raise ValueError(msg)
    if unit == "十人":
        return round(cell * 10)
    return float(cell) if not float(cell).is_integer() else int(cell)
