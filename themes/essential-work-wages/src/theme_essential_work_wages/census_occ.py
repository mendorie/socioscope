"""令和 2 年国勢調査 抽出詳細集計 第 10 表「男女，従業上の地位，職業（小分類）別就業者数（15歳以上）
－全国，都道府県，21大都市，特別区，人口50万以上の市」(e-Stat, 総務省統計局): the national block
(職業 小分類 × 男女 15 歳以上就業者数) -> staged -> mart with ``female_share``.

Table: e-Stat toukei=00200521, tstat=000001136464 (令和 2 年国勢調査), tclass1=000001136468
(抽出詳細集計 主な内容: 就業者の産業・職業（小・中分類）), statInfId ``000032266822``, fileKind=0
(xlsx, 8.3 MB, one sheet ``d10`` with 89,706 rows: 全国 + every prefecture / big city).
Granularity adopted: **小分類** (232 minor groups) plus the 57 中分類 and 12 大分類 rows that come
with them (``level`` column). Classification: 令和 2 年国勢調査に用いる職業分類 (based on 日本標準
職業分類 平成 21 年 12 月改定; ``classification`` = ``jp_census2020``).

Sheet layout (checked 2026-10-29): rows 0-10 title / notes / column headers, row 11 = the header
of the key columns, data from row 12. Columns: 0 地域識別コード (``a``), 1 地域名 (``00000_全国``,
``01000_北海道`` ...), 2 男女 (``0_総数`` / ``1_男`` / ``2_女``), 3 階層レベル（職業） (``1`` major,
``2`` middle, ``3`` minor), 4 職業（大分類）, 5 職業（中分類）, 6 職業 (``<code>_<name>``;
``0_総数`` on the all-occupations row), 7 就業者数 総数 (従業上の地位 0_総数), 8.. 従業上の地位 の
内訳 (not loaded). Counts are sample-based estimates rounded to tens (note 1 of the sheet), so the
sum of the minor groups differs from the national total by rounding; ``-`` means no case.
Only the ``00000_全国`` rows are used (3 sexes × 302 rows: 総数 + 12 majors + 57 middles + 232
minors). Check: for each sex the sum of the minor rows must be within ``TOLERANCE`` (0.1 %) of the
national 総数 row (分類不能の職業 L/99/999 included), else ``ValueError``.

Pure transforms only in ``workbook_rows`` / ``parse`` / ``build_mart``; I/O in fetch / stage / mart
through the ports (see CLAUDE.md 「ソースモジュールの契約」).

License: 政府標準利用規約（第 2.0 版、CC BY 4.0 互換）. 出典「政府統計の総合窓口(e-Stat)、令和 2 年
国勢調査 抽出詳細集計（総務省統計局）を加工」.
"""

from __future__ import annotations

import io
import re
from collections.abc import Sequence

from socioscope_core.core.pipeline import Context
from socioscope_core.ports.fetcher import FetchError
from theme_essential_work_wages.pipeline import THEME

SOURCE = "estat_census2020"
LICENSE = (
    "政府標準利用規約（第2.0版、CC BY 4.0 互換）。出典: 政府統計の総合窓口(e-Stat)、"
    "令和2年国勢調査 抽出詳細集計 第10表（総務省統計局）を加工"
)
STAT_INF_ID = "000032266822"
FILE_KIND = 0
URL = f"https://www.e-stat.go.jp/stat-search/file-download?statInfId={STAT_INF_ID}&fileKind={FILE_KIND}"
RAW_NAME = "census2020_table10_occupation_minor_sex.xlsx"
MAGIC = b"PK\x03\x04"
CENSUS_YEAR = 2020
CLASSIFICATION = "jp_census2020"
STAGED_TABLE = f"staged/{SOURCE}/occupation_sex"
MART_TABLE = "marts/jp_census2020_occupation_sex"
TOLERANCE = 0.001  # relative gap allowed between the national total and the sum of minors

Rows = list[list[object]]
Row = dict[str, object]

COL_AREA, COL_SEX, COL_LEVEL, COL_OCC, COL_TOTAL = 1, 2, 3, 6, 7
NATIONAL = "00000_全国"
SEXES = {"0_総数": "workers_total", "1_男": "workers_male", "2_女": "workers_female"}
LEVELS = {"1": "major", "2": "middle", "3": "minor"}
LEVEL_ORDER = {"major": 0, "middle": 1, "minor": 2}
TOTAL_CODE = "0"
_CODE_NAME = re.compile(r"^([0-9A-Za-z]+)_(.+)$")

MART_COLUMNS = (
    "iso3",
    "year",
    "occupation_code",
    "occupation",
    "level",
    "classification",
    "workers_total",
    "workers_male",
    "workers_female",
    "female_share",
    "source",
)


def workbook_rows(payload: bytes) -> Rows:
    """First sheet of the xlsx as a list of rows (values only)."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as e:
        msg = "payload is not a readable xlsx"
        raise ValueError(msg) from e
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def _text(v: object) -> str:
    return "" if v is None else str(v).replace("　", "").strip()


def _num(v: object) -> float | None:
    """Count cell -> float; '-', 'ｰ', '…', 'X' and blanks -> None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, int | float):
        return float(v)
    s = _text(v).replace(",", "")
    return float(s) if re.match(r"^-?\d+(\.\d+)?$", s) else None


def parse(rows: Sequence[Sequence[object]]) -> list[Row]:
    """National block -> one row per occupation (major / middle / minor) with the 3 sex counts.

    Row order follows the sheet (hierarchical). Raises ``ValueError`` when a sex block or the
    national 総数 row is missing, or when the minors do not add up to the 総数 (±0.1 %).
    """
    counts: dict[str, dict[str, float | None]] = {}  # code -> {sex column -> count}
    meta: dict[str, tuple[str, str]] = {}  # code -> (name, level), in sheet order
    totals: dict[str, float | None] = {}
    for r in rows:
        if len(r) <= COL_TOTAL or _text(r[COL_AREA]) != NATIONAL:
            continue
        sex = SEXES.get(_text(r[COL_SEX]))
        level = LEVELS.get(_text(r[COL_LEVEL]))
        m = _CODE_NAME.match(_text(r[COL_OCC]))
        if sex is None or level is None or m is None:
            continue
        code, name = m.group(1), m.group(2)
        value = _num(r[COL_TOTAL])
        if code == TOTAL_CODE:
            totals[sex] = value
            continue
        meta.setdefault(code, (name, level))
        counts.setdefault(code, {})[sex] = value
    if not meta or not totals:
        msg = "census_occ: no national (00000_全国) occupation rows found"
        raise ValueError(msg)
    for sex_label, sex in SEXES.items():
        if sex not in totals or totals[sex] is None:
            msg = f"census_occ: national 総数 row for {sex_label} missing"
            raise ValueError(msg)
        minors = [counts[c].get(sex) for c, (_, level) in meta.items() if level == "minor"]
        if any(c not in counts or sex not in counts[c] for c in meta):
            msg = f"census_occ: sex block {sex_label} lacks some occupations"
            raise ValueError(msg)
        s = sum(v for v in minors if v is not None)
        total = float(str(totals[sex]))
        if total <= 0 or abs(s - total) / total > TOLERANCE:
            msg = f"census_occ {sex_label}: sum of minors {s} != national total {total}"
            raise ValueError(msg)
    return [
        {
            "occupation_code": code,
            "occupation": name,
            "level": level,
            "workers_total": counts[code]["workers_total"],
            "workers_male": counts[code]["workers_male"],
            "workers_female": counts[code]["workers_female"],
        }
        for code, (name, level) in meta.items()
    ]


def build_mart(staged: Sequence[Row]) -> list[Row]:
    """Add iso3 / year / classification / source and ``female_share``; sort by level then code."""
    out: list[Row] = []
    for r in staged:
        total, female = r.get("workers_total"), r.get("workers_female")
        share = (
            float(str(female)) / float(str(total))
            if isinstance(total, int | float) and isinstance(female, int | float) and total > 0
            else None
        )
        out.append(
            {
                "iso3": "JPN",
                "year": CENSUS_YEAR,
                "occupation_code": r["occupation_code"],
                "occupation": r["occupation"],
                "level": r["level"],
                "classification": CLASSIFICATION,
                "workers_total": total,
                "workers_male": r.get("workers_male"),
                "workers_female": female,
                "female_share": share,
                "source": SOURCE,
            }
        )
    out.sort(key=lambda r: (LEVEL_ORDER.get(str(r["level"]), 9), str(r["occupation_code"])))
    codes = [str(r["occupation_code"]) for r in out]
    if len(set(codes)) != len(codes):
        msg = "census_occ mart: duplicate occupation_code"
        raise ValueError(msg)
    return out


# ---- stages (I/O through ports)


def fetch(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    try:
        payload = ctx.fetcher.fetch(URL)
    except FetchError as e:
        skipped.append(f"{RAW_NAME}: {e}")
        return written, skipped
    if not payload.startswith(MAGIC):
        skipped.append(f"{RAW_NAME}: non-xlsx payload")
        return written, skipped
    rec = ctx.raw.put(
        theme=THEME, source=SOURCE, name=RAW_NAME, url=URL, license=LICENSE, payload=payload
    )
    written.append(rec.relative_path)
    return written, skipped


def stage(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    payload = ctx.raw.get(theme=THEME, source=SOURCE, name=RAW_NAME)
    if payload is None:
        skipped.append(f"{RAW_NAME}: no raw")
        return written, skipped
    ctx.tables.write_table(STAGED_TABLE, parse(workbook_rows(payload)))
    written.append(STAGED_TABLE)
    return written, skipped


def mart(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    try:
        staged = ctx.tables.read_table(STAGED_TABLE)
    except FileNotFoundError:
        skipped.append(f"{STAGED_TABLE}: no staged table")
        return written, skipped
    ctx.tables.write_table(MART_TABLE, build_mart(staged))
    written.append(MART_TABLE)
    return written, skipped
