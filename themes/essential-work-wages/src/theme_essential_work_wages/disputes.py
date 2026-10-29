"""労働争議統計調査 (厚生労働省) 報告書 統計表 第 2 表「産業（大・中分類）、労働争議の種類別件数、
参加人員及び労働損失日数」, calendar years 2012-2024 -> 産業大分類 × 指標 (I/O through ports only).

Tables (e-Stat toukei=00450108, tstat=000001014031, one 報告書/統計表 node per year;
``stat-search/file-download?statInfId=<id>&fileKind=<k>``; all 13 files are xlsx even where
fileKind=0 — checked 2026-10-29): see ``TABLES``. The 2012 table is titled 「産業、労働争議の種類別
…」 (no 「（大・中分類）」) but has the same layout. Industries are 日本標準産業分類 大分類 A-T
(争議統計 uses the 2007/2013 revision; names are kept as in the file, e.g. 「医療，福祉」 with a
full-width comma, and are *not* mapped to the census occupation classification).

Sheet layout (one sheet, 124 rows × 43 cols): rows 0-7 title / 2-level column headers / units,
data from row 8. Column 0 = 大分類 letter (full-width Ａ..Ｔ) or 「全産業」, column 2 = 大分類
name (for 中分類 rows column 2 = 2-digit code and column 3 = name; those are skipped), then
5 総争議 件数, 6 総争議 総参加人員, 7-8 争議行為を伴わない争議 件数/総参加人員, 9 争議行為を伴う
争議 計 件数, 10 同 総参加人員, 11 同 行為参加人員, 12-14 半日以上の同盟罷業及び作業所閉鎖 小計
件数/行為参加人員/労働損失日数, 15-26 行為形態別 (not loaded). ``-`` = no case. 2012-2013 files
store numbers as strings with thousands separators. The only 労働損失日数 in the table is for
半日以上の同盟罷業及び作業所閉鎖 (column 14) — that is ``days_lost``; ``participants`` is the
行為参加人員 of 争議行為を伴う争議 (column 11) and ``participants_total`` the 総参加人員 of 総争議
(column 6). Check: 20 majors present and the 全産業 row equals the sum of the majors for
disputes_total / disputes_with_action / days_lost, else ``ValueError``.

License: 政府標準利用規約（第 2.0 版、CC BY 4.0 互換）.
出典「政府統計の総合窓口(e-Stat)、労働争議統計調査（厚生労働省）を加工」.
"""

from __future__ import annotations

import io
import re
from collections.abc import Sequence

from socioscope_core.core.pipeline import Context
from socioscope_core.ports.fetcher import FetchError
from theme_essential_work_wages.pipeline import THEME

SOURCE = "mhlw_disputes"
LICENSE = (
    "政府標準利用規約（第2.0版、CC BY 4.0 互換）。出典: 政府統計の総合窓口(e-Stat)、"
    "労働争議統計調査 報告書 第2表（厚生労働省）を加工"
)
BASE = "https://www.e-stat.go.jp/stat-search/file-download"
MAGIC = b"PK\x03\x04"
STAGED_PREFIX = f"staged/{SOURCE}/industry_"
MART_TABLE = "marts/jp_labour_disputes_by_industry"

# year -> (statInfId, fileKind). 2012-2019 are listed as fileKind=0 (Excel) but are xlsx too.
TABLES: dict[int, tuple[str, int]] = {
    2012: ("000023598519", 0),
    2013: ("000027341256", 0),
    2014: ("000031340319", 0),
    2015: ("000031478333", 0),
    2016: ("000031643573", 0),
    2017: ("000031768817", 0),
    2018: ("000031877549", 0),
    2019: ("000032018193", 0),
    2020: ("000032128457", 4),
    2021: ("000032250878", 4),
    2022: ("000040119703", 4),
    2023: ("000040225302", 4),
    2024: ("000040372369", 4),
}

Rows = list[list[object]]
Row = dict[str, object]

COL_CODE, COL_NAME = 0, 2
COL_TOTAL, COL_PART_TOTAL, COL_WITH_ACTION, COL_PART_ACTION, COL_DAYS = 5, 6, 9, 11, 14
TOTAL_LABEL = "全産業"
TOTAL_CODE = "ALL"
MAJORS = "ABCDEFGHIJKLMNOPQRST"
_FULLWIDTH_A = ord("Ａ")
METRICS = (
    "disputes_total",
    "participants_total",
    "disputes_with_action",
    "participants",
    "days_lost",
)
CHECKED = ("disputes_total", "disputes_with_action", "days_lost")
MART_COLUMNS = ("iso3", "year", "industry_code", "industry", *METRICS, "source")


def url(year: int) -> str:
    sid, kind = TABLES[year]
    return f"{BASE}?statInfId={sid}&fileKind={kind}"


def raw_name(year: int) -> str:
    return f"disputes_{year}_industry.xlsx"


def staged_table(year: int) -> str:
    return f"{STAGED_PREFIX}{year}"


def workbook_rows(payload: bytes) -> Rows:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as e:
        msg = "payload is not a readable xlsx"
        raise ValueError(msg) from e
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def _text(v: object) -> str:
    return "" if v is None else str(v).replace("　", "").strip()


def _num(v: object) -> float | None:
    """Cell -> float; '-', 'ｰ', '…', 'X' and blanks -> None (strings may carry commas)."""
    if isinstance(v, bool):
        return None
    if isinstance(v, int | float):
        return float(v)
    s = _text(v).replace(",", "")
    return float(s) if re.match(r"^-?\d+(\.\d+)?$", s) else None


def _major_code(v: object) -> str | None:
    s = _text(v)
    if len(s) != 1:
        return None
    if "Ａ" <= s <= "Ｔ":
        return chr(ord(s) - _FULLWIDTH_A + ord("A"))
    return s if s in MAJORS else None


def _metrics(r: Sequence[object]) -> dict[str, float | None]:
    return {
        "disputes_total": _num(r[COL_TOTAL]),
        "participants_total": _num(r[COL_PART_TOTAL]),
        "disputes_with_action": _num(r[COL_WITH_ACTION]),
        "participants": _num(r[COL_PART_ACTION]),
        "days_lost": _num(r[COL_DAYS]),
    }


def parse(rows: Sequence[Sequence[object]], year: int) -> list[Row]:
    """Sheet rows -> [全産業 row, then majors A..T] for ``year``. Fail-closed on layout drift."""
    out: list[Row] = []
    total: Row | None = None
    for r in rows:
        if len(r) <= COL_DAYS:
            continue
        if _text(r[COL_CODE]) == TOTAL_LABEL and total is None:
            total = {
                "year": year,
                "industry_code": TOTAL_CODE,
                "industry": TOTAL_LABEL,
                **_metrics(r),
            }
            continue
        code = _major_code(r[COL_CODE])
        name = _text(r[COL_NAME])
        if code is None or not name:
            continue
        out.append({"year": year, "industry_code": code, "industry": name, **_metrics(r)})
    if total is None:
        msg = f"disputes {year}: 全産業 row missing"
        raise ValueError(msg)
    codes = [str(r["industry_code"]) for r in out]
    if codes != list(MAJORS):
        msg = f"disputes {year}: expected 20 major industries A..T, got {codes}"
        raise ValueError(msg)
    for metric in CHECKED:
        s = sum(float(str(r[metric])) for r in out if r[metric] is not None)
        t = total[metric]
        if t is None or abs(float(str(t)) - s) > 0.5:
            msg = f"disputes {year}: {metric} majors sum {s} != 全産業 {t}"
            raise ValueError(msg)
    return [total, *out]


def build_mart(staged: Sequence[Row]) -> list[Row]:
    out: list[Row] = [
        {
            "iso3": "JPN",
            "year": int(str(r["year"])),
            "industry_code": r["industry_code"],
            "industry": r["industry"],
            **{m: r.get(m) for m in METRICS},
            "source": SOURCE,
        }
        for r in staged
    ]
    out.sort(key=lambda r: (int(str(r["year"])), str(r["industry_code"])))
    seen: set[tuple[int, str]] = set()
    for r in out:
        k = (int(str(r["year"])), str(r["industry_code"]))
        if k in seen:
            msg = f"disputes mart: duplicate year/industry {k}"
            raise ValueError(msg)
        seen.add(k)
    return out


# ---- stages (I/O through ports)


def fetch(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    for year in sorted(TABLES):
        name, u = raw_name(year), url(year)
        try:
            payload = ctx.fetcher.fetch(u)
        except FetchError as e:
            skipped.append(f"{name}: {e}")
            continue
        if not payload.startswith(MAGIC):
            skipped.append(f"{name}: non-xlsx payload")
            continue
        rec = ctx.raw.put(
            theme=THEME, source=SOURCE, name=name, url=u, license=LICENSE, payload=payload
        )
        written.append(rec.relative_path)
    return written, skipped


def stage(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    for year in sorted(TABLES):
        payload = ctx.raw.get(theme=THEME, source=SOURCE, name=raw_name(year))
        if payload is None:
            skipped.append(f"{raw_name(year)}: no raw")
            continue
        ctx.tables.write_table(staged_table(year), parse(workbook_rows(payload), year))
        written.append(staged_table(year))
    return written, skipped


def mart(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    staged: list[Row] = []
    for year in sorted(TABLES):
        try:
            staged.extend(ctx.tables.read_table(staged_table(year)))
        except FileNotFoundError:
            continue
    if not staged:
        skipped.append(f"{SOURCE}: no staged table")
        return written, skipped
    ctx.tables.write_table(MART_TABLE, build_mart(staged))
    written.append(MART_TABLE)
    return written, skipped
