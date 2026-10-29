"""財務省 法人企業統計調査 年次別調査「2. 業種別、規模別資産・負債・純資産及び損益表」
（e-Stat、appId 不要）。

H6 のレント指標 R（産業の 1 人当たり付加価値・営業利益率、日本）に使う。
e-Stat 政府統計コード 00350600（`tstat=000001047744`、年度次 > 年次別調査 > 各年度 > 統計表 2）。
年度ごとに 1 ブック（`file-download?statInfId=<STAT_INF_IDS[fy]>&fileKind=0`、約 0.7〜1.2 MB）、
FY2012〜FY2018 は BIFF `.xls`（xlrd）、FY2019〜FY2022 は `.xlsx`（openpyxl）。
レイアウトは全年度同じ:
1 シート = 1 業種（68 シート。`全産業` → 製造業とその中分類 → `非製造業` とその大・中分類 →
`全産業（金融業、保険業を含む）` 以降は金融・保険の別様式）。行 = 項目（列 1〜6 に階層付き項目名）、
列 = 資本金階級（列 7 総額、列 8〜12 1,000 万円未満〜10 億円以上）。単位は百万円・人。
負数は FY2013 以前は '△ 30,070' の文字列、FY2014 以降は数値。

採る項目: 売上高、営業利益、付加価値額、人件費（役員給与＋役員賞与＋従業員給与＋従業員賞与＋
福利厚生費。法人企業統計の「人件費」の定義）、従業員数（人）。金融・保険のシートは売上高・営業利益・
付加価値額を持たない別様式なので読まない（`全産業（金融業、保険業を含む）` のシートで打ち切る）。
`industry_code` は本モジュールが与える安定コード（`INDUSTRIES`。JSIC 大分類の英字に寄せた
記号で、公式コードではない）。`industry_level`: total / aggregate / major / minor。

検算: 総額列で 全産業 ＝ 製造業 ＋ 非製造業 が売上高・付加価値額・従業員数について ±0.5% に
収まらなければ ValueError。想定外のシート名・項目行欠落・数値でないセル（'－' 以外）も ValueError。

FY2023・FY2024 の年次別調査は e-Stat に xlsx が無く（報道発表 PDF と時系列 DB のみ、2026-10-29
確認）、本モジュールの対象外。追加するときは `STAT_INF_IDS` に足す。

License: 政府標準利用規約（第 2.0 版）＝ CC BY 4.0 互換（出典: 財務省「法人企業統計調査」）。
"""

from __future__ import annotations

import io
import re
from collections.abc import Mapping, Sequence

from socioscope_core.core.pipeline import Context
from socioscope_core.ports.fetcher import FetchError
from theme_essential_work_wages.pipeline import THEME

SOURCE = "mof_hojin"
LICENSE = (
    "政府標準利用規約（第2.0版）/ CC BY 4.0 compatible (財務省 法人企業統計調査, e-Stat 00350600)"
)
BASE = "https://www.e-stat.go.jp/stat-search/file-download"
MAGIC: tuple[bytes, ...] = (b"PK", b"\xd0\xcf\x11\xe0")  # xlsx (zip) / BIFF xls
MART_TABLE = "marts/hojin_industry_annual"
TOLERANCE = 0.005
# 年度 -> 統計表 2 の statInfId（e-Stat 年次別調査、2026-10-29 確認）
STAT_INF_IDS: dict[int, str] = {
    2012: "000023757923",
    2013: "000030052839",
    2014: "000031369065",
    2015: "000031450735",
    2016: "000031621132",
    2017: "000031751513",
    2018: "000031860693",
    2019: "000032015945",
    2020: "000032192030",
    2021: "000040090169",
    2022: "000040172276",
}
# シート名（正規化後）-> (code, parent_code, level)。並びは実表と同じ。
INDUSTRIES: dict[str, tuple[str, str | None, str]] = {
    "全産業": ("ALL_EXFIN", None, "total"),
    "製造業": ("E", "ALL_EXFIN", "major"),
    "食料品製造業": ("E01", "E", "minor"),
    "繊維工業": ("E02", "E", "minor"),
    "木材・木製品製造業": ("E03", "E", "minor"),
    "パルプ・紙・紙加工品製造業": ("E04", "E", "minor"),
    "印刷・同関連業": ("E05", "E", "minor"),
    "化学工業": ("E06", "E", "minor"),
    "石油製品・石炭製品製造業": ("E07", "E", "minor"),
    "窯業・土石製品製造業": ("E08", "E", "minor"),
    "鉄鋼業": ("E09", "E", "minor"),
    "非鉄金属製造業": ("E10", "E", "minor"),
    "金属製品製造業": ("E11", "E", "minor"),
    "はん用機械器具製造業": ("E12", "E", "minor"),
    "生産用機械器具製造業": ("E13", "E", "minor"),
    "業務用機械器具製造業": ("E14", "E", "minor"),
    "電気機械器具製造業": ("E15", "E", "minor"),
    "情報通信機械器具製造業": ("E16", "E", "minor"),
    "輸送用機械器具製造業": ("E17", "E", "minor"),
    "自動車・同附属品製造業": ("E17a", "E17", "minor"),
    "その他の輸送用機械器具製造業": ("E17b", "E17", "minor"),
    "その他の製造業": ("E18", "E", "minor"),
    "非製造業": ("NONMFG", "ALL_EXFIN", "aggregate"),
    "農林水産業": ("AB", "NONMFG", "major"),
    "農業、林業": ("A", "AB", "minor"),
    "漁業": ("B", "AB", "minor"),
    "鉱業、採石業、砂利採取業": ("C", "NONMFG", "major"),
    "建設業": ("D", "NONMFG", "major"),
    "電気業": ("F1", "NONMFG", "major"),
    "ガス・熱供給・水道業": ("F2", "NONMFG", "major"),
    "情報通信業": ("G", "NONMFG", "major"),
    "運輸業、郵便業": ("H", "NONMFG", "major"),
    "陸運業": ("H1", "H", "minor"),
    "水運業": ("H2", "H", "minor"),
    "その他の運輸業": ("H3", "H", "minor"),
    "卸売業、小売業": ("I", "NONMFG", "major"),
    "卸売業": ("I1", "I", "minor"),
    "小売業": ("I2", "I", "minor"),
    "不動産業、物品賃貸業": ("K", "NONMFG", "major"),
    "不動産業": ("K1", "K", "minor"),
    "物品賃貸業": ("K2", "K", "minor"),
    "リース業": ("K2a", "K2", "minor"),
    "その他の物品賃貸業": ("K2b", "K2", "minor"),
    "サービス業": ("SVC", "NONMFG", "aggregate"),
    "宿泊業、飲食サービス業": ("M", "SVC", "major"),
    "宿泊業": ("M1", "M", "minor"),
    "飲食サービス業": ("M2", "M", "minor"),
    "生活関連サービス業、娯楽業": ("N", "SVC", "major"),
    "生活関連サービス業": ("N1", "N", "minor"),
    "娯楽業": ("N2", "N", "minor"),
    "学術研究、専門・技術サービス業": ("L", "SVC", "major"),
    "広告業": ("L1", "L", "minor"),
    "純粋持株会社": ("L2", "L", "minor"),
    "その他の学術研究、専門・技術サービス業": ("L3", "L", "minor"),
    "医療、福祉業": ("P", "SVC", "major"),
    "教育、学習支援業": ("O", "SVC", "major"),
    "職業紹介・労働者派遣業": ("R1", "SVC", "major"),
    "その他のサービス業": ("R2", "SVC", "major"),
}
FINANCE_MARKER = "全産業(金融業、保険業を含む)"  # このシート以降は別様式。読まない
SHEET_ALIASES = {"ガス・熱供給水道業": "ガス・熱供給・水道業"}  # FY2012 の表記ゆれ
# 項目名（正規化後）-> 出力列。人件費は 5 項目の和
ITEMS: dict[str, str] = {
    "売上高": "sales_mn_yen",
    "営業利益": "operating_profit_mn_yen",
    "付加価値額": "value_added_mn_yen",
    "従業員数(人)": "employees",
}
PERSONNEL_ITEMS = ("役員給与", "役員賞与", "従業員給与", "従業員賞与", "福利厚生費")
SIZE_CLASSES: dict[str, str] = {  # 資本金（百万円）階級の見出し -> size_class
    "総額": "all",
    "10未満": "lt_10",
    "10～50未満": "10_50",
    "50～100未満": "50_100",
    "100～1,000未満": "100_1000",
    "1,000以上": "ge_1000",
}
CHECKED = ("sales_mn_yen", "value_added_mn_yen", "employees")
LABEL_COLS = range(0, 7)
Rows = list[list[object]]


def url(fiscal_year: int) -> str:
    return f"{BASE}?statInfId={STAT_INF_IDS[fiscal_year]}&fileKind=0"


def raw_name(fiscal_year: int) -> str:
    return f"hojin_t2_fy{fiscal_year}.xls"


def staged_table(fiscal_year: int) -> str:
    return f"staged/{SOURCE}/hojin_t2_fy{fiscal_year}"


def workbook_rows(payload: bytes) -> list[tuple[str, Rows]]:
    """xlsx / xls → [(シート名, 行リスト)]。読めなければ ValueError。"""
    try:
        if payload.startswith(MAGIC[0]):
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            return [
                (ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
                for ws in wb.worksheets
            ]
        import xlrd

        bk = xlrd.open_workbook(file_contents=payload)
        return [
            (sh.name, [[_blank_to_none(c) for c in sh.row_values(i)] for i in range(sh.nrows)])
            for sh in bk.sheets()
        ]
    except Exception as e:
        msg = f"hojin: payload is not a readable xls/xlsx workbook ({type(e).__name__})"
        raise ValueError(msg) from e


def _blank_to_none(v: object) -> object:
    return None if v == "" else v


def _norm(v: object) -> str:
    s = "" if v is None else str(v)
    s = re.sub(r"[\s　]+", "", s)
    return s.replace("（", "(").replace("）", ")").replace("･", "・")


def _num(v: object, where: str) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int | float):
        return float(v)
    s = _norm(v).replace(",", "")
    if s in ("－", "-", "―", "…", ""):
        return None
    s = re.sub(r"^[△▲]", "-", s)  # FY2013 以前の負数表記（'△ 30,070'）
    if re.match(r"^-?\d+(\.\d+)?$", s):
        return float(s)
    msg = f"hojin: not a number at {where}: {v!r}"
    raise ValueError(msg)


def _label(row: list[object]) -> str:
    return next((_norm(c) for c in row[: LABEL_COLS.stop] if c not in (None, "")), "")


def _size_columns(name: str, rows: Rows) -> list[tuple[int, str]]:
    for row in rows[:8]:
        cells = {_norm(c): i for i, c in enumerate(row) if c is not None}
        if all(h in cells for h in SIZE_CLASSES):
            return [(cells[h], sc) for h, sc in SIZE_CLASSES.items()]
    msg = f"hojin: sheet {name!r} has no size header {tuple(SIZE_CLASSES)}"
    raise ValueError(msg)


def _item_rows(name: str, rows: Rows) -> dict[str, list[object]]:
    wanted = set(ITEMS) | set(PERSONNEL_ITEMS)
    found: dict[str, list[object]] = {}
    for row in rows:
        lab = _label(row)
        if lab in wanted and lab not in found:
            found[lab] = row
    missing = [w for w in (*ITEMS, *PERSONNEL_ITEMS) if w not in found]
    if missing:
        msg = f"hojin: sheet {name!r} lacks item rows {missing}"
        raise ValueError(msg)
    return found


def _parse_sheet(
    title: str, rows: Rows, fiscal_year: int, industry: str
) -> list[dict[str, object]]:
    code, parent, level = INDUSTRIES[industry]
    items = _item_rows(title, rows)
    out: list[dict[str, object]] = []
    for col, size_class in _size_columns(title, rows):
        rec: dict[str, object] = {
            "fiscal_year": fiscal_year,
            "industry": industry,
            "industry_code": code,
            "industry_level": level,
            "parent_code": parent,
            "size_class": size_class,
        }
        for item, column in ITEMS.items():
            v = _num(_cell(items[item], col), f"{title}/{item}/{size_class}")
            rec[column] = int(v) if column == "employees" and v is not None else v
        parts = [_num(_cell(items[p], col), f"{title}/{p}/{size_class}") for p in PERSONNEL_ITEMS]
        rec["personnel_cost_mn_yen"] = (
            sum(p for p in parts if p is not None) if any(p is not None for p in parts) else None
        )
        # 列順を固定（spec の並び）
        out.append(
            {
                **{k: rec[k] for k in ("fiscal_year", "industry", "industry_code")},
                "industry_level": rec["industry_level"],
                "parent_code": rec["parent_code"],
                "size_class": rec["size_class"],
                "sales_mn_yen": rec["sales_mn_yen"],
                "operating_profit_mn_yen": rec["operating_profit_mn_yen"],
                "personnel_cost_mn_yen": rec["personnel_cost_mn_yen"],
                "value_added_mn_yen": rec["value_added_mn_yen"],
                "employees": rec["employees"],
            }
        )
    return out


def _cell(row: list[object], col: int) -> object:
    return row[col] if col < len(row) else None


def parse(sheets: Sequence[tuple[str, Rows]], fiscal_year: int) -> list[dict[str, object]]:
    """ブック（[(シート名, 行)]）→ 業種 × 資本金階級 の行。

    総額列で 全産業＝製造業＋非製造業 を検算（±0.5%）。"""
    if not sheets:
        msg = "hojin: workbook has no sheet"
        raise ValueError(msg)
    out: list[dict[str, object]] = []
    for title, rows in sheets:
        name = _norm(title)
        name = SHEET_ALIASES.get(name, name)
        if name == FINANCE_MARKER:
            break
        if name not in INDUSTRIES:
            msg = f"hojin: unknown sheet {title!r} (FY{fiscal_year})"
            raise ValueError(msg)
        out.extend(_parse_sheet(title, rows, fiscal_year, name))
    _check_totals(out, fiscal_year)
    return out


def _check_totals(rows: Sequence[Mapping[str, object]], fiscal_year: int) -> None:
    totals = {
        str(r["industry_code"]): r
        for r in rows
        if r["size_class"] == "all" and r["industry_code"] in ("ALL_EXFIN", "E", "NONMFG")
    }
    missing = [c for c in ("ALL_EXFIN", "E", "NONMFG") if c not in totals]
    if missing:
        names = {"ALL_EXFIN": "全産業", "E": "製造業", "NONMFG": "非製造業"}
        lacking = [names[c] for c in missing]
        msg = f"hojin: FY{fiscal_year} lacks sheet(s) {lacking} for the total check"
        raise ValueError(msg)
    for column in CHECKED:
        whole = totals["ALL_EXFIN"][column]
        parts = (totals["E"][column], totals["NONMFG"][column])
        if whole is None or any(p is None for p in parts):
            msg = f"hojin: FY{fiscal_year} {column} missing in a total row"
            raise ValueError(msg)
        total = float(str(parts[0])) + float(str(parts[1]))
        if abs(float(str(whole)) - total) > TOLERANCE * abs(float(str(whole))):
            msg = (
                f"hojin: FY{fiscal_year} {column}: 全産業 {whole} != 製造業+非製造業 {total} "
                f"(tolerance {TOLERANCE:.1%})"
            )
            raise ValueError(msg)


def _ratio(num: object, den: object, scale: float = 1.0) -> float | None:
    if num is None or den is None or float(str(den)) == 0:
        return None
    return float(str(num)) / float(str(den)) * scale


def _key(r: Mapping[str, object]) -> tuple[int, str, str]:
    return (int(str(r["fiscal_year"])), str(r["industry_code"]), str(r["size_class"]))


def build_mart(staged: Sequence[Mapping[str, object]]) -> list[dict[str, object]]:
    """年度 × 業種 × 規模 の mart。1 人当たり付加価値・営業利益率を付与。キー重複は ValueError。"""
    seen: set[tuple[int, str, str]] = set()
    out: list[dict[str, object]] = []
    for r in staged:
        key = _key(r)
        if key in seen:
            msg = f"hojin: duplicate {key}"
            raise ValueError(msg)
        seen.add(key)
        out.append(
            {
                "iso3": "JPN",
                "year": key[0],
                **dict(r),
                "value_added_per_employee_mn_yen": _ratio(r["value_added_mn_yen"], r["employees"]),
                "operating_margin_pct": _ratio(
                    r["operating_profit_mn_yen"], r["sales_mn_yen"], 100.0
                ),
                "source": SOURCE,
            }
        )
    out.sort(key=_key)
    return out


# ---- stage functions (I/O through ports only)


def fetch(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    for fy in sorted(STAT_INF_IDS):
        name, u = raw_name(fy), url(fy)
        try:
            payload = ctx.fetcher.fetch(u)
        except FetchError as e:
            skipped.append(f"{name}: {e}")
            continue
        if not payload.startswith(MAGIC):
            skipped.append(f"{name}: non-Excel payload")
            continue
        rec = ctx.raw.put(
            theme=THEME, source=SOURCE, name=name, url=u, license=LICENSE, payload=payload
        )
        written.append(rec.relative_path)
    return written, skipped


def stage(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    for fy in sorted(STAT_INF_IDS):
        name = raw_name(fy)
        payload = ctx.raw.get(theme=THEME, source=SOURCE, name=name)
        if payload is None:
            skipped.append(f"{name}: no raw")
            continue
        table = staged_table(fy)
        ctx.tables.write_table(table, parse(workbook_rows(payload), fy))
        written.append(table)
    return written, skipped


def mart(ctx: Context, written: list[str], skipped: list[str]) -> tuple[list[str], list[str]]:
    staged: list[dict[str, object]] = []
    for fy in sorted(STAT_INF_IDS):
        try:
            staged.extend(ctx.tables.read_table(staged_table(fy)))
        except FileNotFoundError:
            continue
    if not staged:
        skipped.append(f"{MART_TABLE}: no staged table")
        return written, skipped
    ctx.tables.write_table(MART_TABLE, build_mart(staged))
    written.append(MART_TABLE)
    return written, skipped
