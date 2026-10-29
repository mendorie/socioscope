"""Census ACS 1-year B24121-B24126 (occupation x median earnings / workers) source module."""

import io
from pathlib import Path

import openpyxl
import pytest

from socioscope_core.core.pipeline import Context
from socioscope_core.testing.fakes import FakeFetcher, InMemoryRawStore, InMemoryTableStore
from theme_essential_work_wages import acs
from theme_essential_work_wages.pipeline import THEME

FIX = Path(__file__).parent / "fixtures"

SUMMARY_ROWS: list[tuple[object, ...]] = [
    (None, None, "2018 Census Occupation Code List - Summary of Changes "),
    (None, None, "Posted September 26,2019 "),
    ("2018 SOC Code", "2018 Census Code", "2018 Census Title ", "Census No Change"),
    ("11-1011", "0010", "Chief Executives", "x"),
    ("11-1021", "0020", "General and Operations Managers", "x"),
    ("11-1021", "0020", "General and Operations Managers", "x"),  # real file repeats 0650
    ("11-1031", "0030", "Legislators", "x"),
    ("55-1010", "9800", "Military Officer Special and Tactical Operations Leaders", "x"),
    ("none", "9830", "Military, Rank Not Specified", "x"),
    (None, None, None, None),
    (None, "Legend", "Yellow fill- new occupation, code, title", None),
]

XWALK_ROWS: list[tuple[object, ...]] = [
    ("Census Occupation Code Crosswalk from 2010 to 2018 ",),
    ("September 26, 2019",),
    (None,),
    (
        "2010 SOC code",
        "2010 Census Code",
        "2010 Census Title \n",
        "2018 SOC Code",
        "2018 Census Code",
        "2018 Census Title ",
    ),
    ("11-1011", "0010", "Chief Executives", "11-1011", "0010", "Chief Executives"),
    ("11-1020", "0020", "Operations Managers (2010)", None, None, None),  # split block source
    ("11-1029", "0021", "Other Managers (2010)", None, None, "merged (11-1029)"),
    (None, None, None, "11-1021", "0020", "General and Operations Managers"),
    ("11-1031", "0030", "Legislators", "11-1031", "0030", "Legislators"),
    (None, None, None, None, None, "Legislators, note (11-1031)"),
    (
        "none",
        "9830",
        "Military, Rank Not Specified",
        "none",
        "9830",
        "Military, Rank Not Specified",
    ),
    (None, None, "Total Codes, 2010: 540 including military", None, None, None),
]


def make_xlsx(
    summary: list[tuple[object, ...]] = SUMMARY_ROWS, xwalk: list[tuple[object, ...]] = XWALK_ROWS
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "OVERVIEW"
    s = wb.create_sheet("Summary of 2018 Changes")
    for r in summary:
        s.append(list(r))
    x = wb.create_sheet("2010 to 2018 Crosswalk ")
    for r in xwalk:
        x.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def dat(table: str) -> bytes:
    return (FIX / f"acsdt1y2023-{table.lower()}_head.dat").read_bytes()


def shells() -> bytes:
    return (FIX / "ACS20231YR_Table_Shells_head.txt").read_bytes()


def responses(years: tuple[int, ...] = acs.YEARS) -> dict[str, bytes]:
    out = {acs.url(y, t): dat(t) for y in years for t in acs.TABLES}
    out[acs.SHELLS_URL] = shells()
    out[acs.CODES_URL] = make_xlsx()
    return out


def make_ctx(resp: dict[str, bytes] | None = None) -> Context:
    return Context(
        fetcher=FakeFetcher(resp or {}), raw=InMemoryRawStore(), tables=InMemoryTableStore()
    )


# ---- constants


def test_urls_and_raw_names() -> None:
    assert acs.SOURCE == "census_acs1"
    assert acs.YEARS == (2021, 2022, 2023, 2024)
    assert acs.TABLES == ("B24121", "B24122", "B24123", "B24124", "B24125", "B24126")
    assert acs.url(2023, "B24121") == (
        "https://www2.census.gov/programs-surveys/acs/summary_file/2023/table-based-SF/"
        "data/1YRData/acsdt1y2023-b24121.dat"
    )
    assert acs.raw_name(2023, "B24121") == "acsdt1y2023-b24121.dat"
    assert acs.SHELLS_URL == (
        "https://www2.census.gov/programs-surveys/acs/summary_file/2023/table-based-SF/"
        "documentation/ACS20231YR_Table_Shells.txt"
    )
    assert acs.SHELLS_RAW_NAME == "ACS20231YR_Table_Shells.txt"
    assert acs.CODES_URL == (
        "https://www2.census.gov/programs-surveys/demo/guidance/industry-occupation/"
        "2018-occupation-code-list-and-crosswalk.xlsx"
    )
    assert acs.CODES_RAW_NAME == "2018-occupation-code-list-and-crosswalk.xlsx"
    assert acs.MAGIC == b"GEO_ID|"
    assert acs.SHELLS_MAGIC == b"Table ID|"
    assert acs.XLSX_MAGIC == b"PK\x03\x04"


# ---- fetch


def test_fetch_puts_every_table_year_plus_shells_and_codes() -> None:
    ctx = make_ctx(responses())
    written, skipped = acs.fetch(ctx, [], [])
    assert skipped == []
    assert len(written) == 4 * 6 + 2
    assert f"{THEME}/{acs.SOURCE}/acsdt1y2021-b24126.dat" in written
    assert f"{THEME}/{acs.SOURCE}/{acs.SHELLS_RAW_NAME}" in written
    assert f"{THEME}/{acs.SOURCE}/{acs.CODES_RAW_NAME}" in written
    recs = ctx.raw.records()
    assert len(recs) == 26
    assert {r.license for r in recs} == {acs.LICENSE}
    assert {r.theme for r in recs} == {THEME}
    assert ctx.raw.get(theme=THEME, source=acs.SOURCE, name="acsdt1y2024-b24121.dat") == dat(
        "B24121"
    )


def test_fetch_reports_fetch_errors_and_rejects_non_magic_payloads() -> None:
    resp = responses()
    del resp[acs.url(2022, "B24123")]
    resp[acs.url(2021, "B24121")] = b"<html>blocked</html>"
    resp[acs.SHELLS_URL] = b"<!DOCTYPE html>"
    resp[acs.CODES_URL] = b"not a zip"
    ctx = make_ctx(resp)
    written, skipped = acs.fetch(ctx, [], [])
    assert len(written) == 24 - 2
    assert any(s.startswith("acsdt1y2022-b24123.dat: ") for s in skipped)
    assert "acsdt1y2021-b24121.dat: non-dat payload" in skipped
    assert f"{acs.SHELLS_RAW_NAME}: non-shells payload" in skipped
    assert f"{acs.CODES_RAW_NAME}: non-xlsx payload" in skipped
    assert ctx.raw.get(theme=THEME, source=acs.SOURCE, name="acsdt1y2021-b24121.dat") is None
    assert ctx.raw.get(theme=THEME, source=acs.SOURCE, name=acs.SHELLS_RAW_NAME) is None
    assert len(ctx.raw.records()) == 22


def test_fetch_accepts_shells_with_utf8_bom() -> None:
    # The 2022 shells file starts with a BOM, 2021/2024 do not.
    resp = responses()
    resp[acs.SHELLS_URL] = b"\xef\xbb\xbf" + shells().lstrip(b"\xef\xbb\xbf")
    ctx = make_ctx(resp)
    _, skipped = acs.fetch(ctx, [], [])
    assert skipped == []


# ---- pure parsers


def test_parse_dat_returns_estimate_and_moe_per_line() -> None:
    out = acs.parse_dat(dat("B24121"), 2023, "B24121")
    assert out == {
        1: (60192.0, 57.0),
        2: (170826.0, 3213.0),
        3: (84493.0, 1344.0),
        4: (69328.0, 5703.0),
    }


def test_parse_dat_turns_census_special_values_into_none() -> None:
    payload = (
        b"GEO_ID|B24123_E001|B24123_M001|B24123_E002|B24123_M002|B24123_E003|B24123_M003\n"
        b"0100000US|-666666666|-222222222|52458|-333333333||(X)\n"
    )
    assert acs.parse_dat(payload, 2023, "B24123") == {
        1: (None, None),
        2: (52458.0, None),
        3: (None, None),
    }


@pytest.mark.parametrize(
    "payload",
    [
        b"GEO_ID|B24122_E001|B24122_M001\n0100000US|1|2\n",  # wrong table
        b"GEO_ID|B24121_E001|B24121_M001\n0400000US01|1|2\n",  # not the national row
        b"GEO_ID|B24121_E001|B24121_M001\n0100000US|1|2\n0100000US|1|2\n",  # 2 data rows
        b"GEO_ID|B24121_E001|B24121_E002\n0100000US|1|2\n",  # E without its M
        b"GEO_ID|B24121_E001|B24121_M001\n0100000US|1\n",  # ragged
        b"GEO_ID|B24121_E001|B24121_M001\n0100000US|abc|2\n",  # non-numeric
    ],
)
def test_parse_dat_rejects_unexpected_layout(payload: bytes) -> None:
    with pytest.raises(ValueError):
        acs.parse_dat(payload, 2023, "B24121")


def test_parse_shells_picks_the_table_lines_in_order() -> None:
    assert acs.parse_shells(shells(), "B24124") == [
        (1, "Total:"),
        (2, "Chief executives"),
        (3, "General and operations managers"),
        (4, "Legislators"),
    ]


def test_parse_shells_rejects_missing_table() -> None:
    with pytest.raises(ValueError):
        acs.parse_shells(shells(), "B24999")


def test_parse_codes_reads_summary_sheet_sorted_and_deduplicated() -> None:
    assert acs.parse_codes(make_xlsx()) == [
        ("0010", "11-1011", "Chief Executives"),
        ("0020", "11-1021", "General and Operations Managers"),
        ("0030", "11-1031", "Legislators"),
        ("9800", "55-1010", "Military Officer Special and Tactical Operations Leaders"),
        ("9830", "none", "Military, Rank Not Specified"),
    ]


def test_parse_codes_rejects_conflicting_duplicate_codes() -> None:
    rows = [*SUMMARY_ROWS, ("11-9999", "0010", "Someone Else", None)]
    with pytest.raises(ValueError, match="0010"):
        acs.parse_codes(make_xlsx(summary=rows))


def test_parse_crosswalk_maps_2018_codes_to_their_2010_soc_codes() -> None:
    assert acs.parse_crosswalk(make_xlsx()) == {
        "0010": ("11-1011",),
        "0020": ("11-1020", "11-1029"),
        "0030": ("11-1031",),
        "9830": ("none",),
    }


def test_build_codes_table_keeps_civilian_codes_with_joined_2010_socs() -> None:
    rows = acs.build_codes_table(acs.parse_codes(make_xlsx()), acs.parse_crosswalk(make_xlsx()))
    assert rows == [
        {
            "occupation_code": "0010",
            "occupation": "Chief Executives",
            "soc_code": "11-1011",
            "soc_2010_codes": "11-1011",
            "classification": acs.CLASSIFICATION,
        },
        {
            "occupation_code": "0020",
            "occupation": "General and Operations Managers",
            "soc_code": "11-1021",
            "soc_2010_codes": "11-1020;11-1029",
            "classification": acs.CLASSIFICATION,
        },
        {
            "occupation_code": "0030",
            "occupation": "Legislators",
            "soc_code": "11-1031",
            "soc_2010_codes": "11-1031",
            "classification": acs.CLASSIFICATION,
        },
    ]


def test_build_codes_table_falls_back_to_empty_string_when_code_is_not_in_crosswalk() -> None:
    # Real file: 8225 "Other Metal Workers and Plastic Workers" (51-4XXX) is new in 2018.
    codes = [("8225", "51-4XXX", "Other Metal Workers and Plastic Workers")]
    assert acs.build_codes_table(codes, {})[0]["soc_2010_codes"] == ""


# ---- stage


def test_stage_reports_no_raw_and_writes_nothing() -> None:
    ctx = make_ctx()
    written, skipped = acs.stage(ctx, [], [])
    assert written == []
    assert f"{acs.SHELLS_RAW_NAME}: no raw" in skipped
    assert f"{acs.CODES_RAW_NAME}: no raw" in skipped
    assert "acsdt1y2021-b24121.dat: no raw" in skipped
    assert ctx.tables.list_tables() == []


def test_stage_skips_a_year_when_one_of_its_tables_is_missing() -> None:
    resp = responses()
    del resp[acs.url(2022, "B24125")]
    ctx = make_ctx(resp)
    acs.fetch(ctx, [], [])
    written, skipped = acs.stage(ctx, [], [])
    assert acs.staged_table(2022) not in written
    assert acs.staged_table(2021) in written
    assert skipped == ["acsdt1y2022-b24125.dat: no raw"]


def test_stage_writes_one_wide_table_per_year_and_the_codes_table() -> None:
    ctx = make_ctx(responses())
    acs.fetch(ctx, [], [])
    written, skipped = acs.stage(ctx, [], [])
    assert skipped == []
    assert written == [acs.CODES_TABLE, *[acs.staged_table(y) for y in acs.YEARS]]
    rows = ctx.tables.read_table(acs.staged_table(2023))
    assert len(rows) == 4
    assert rows[1] == {
        "year": 2023,
        "line": 2,
        "label": "Chief executives",
        "b24121_e": 170826.0,
        "b24121_m": 3213.0,
        "b24122_e": 184864.0,
        "b24122_m": 3842.0,
        "b24123_e": 145497.0,
        "b24123_m": 4464.0,
        "b24124_e": 1439966.0,
        "b24124_m": 19992.0,
        "b24125_e": 1009967.0,
        "b24125_m": 16858.0,
        "b24126_e": 429999.0,
        "b24126_m": 10538.0,
    }
    assert [r["occupation_code"] for r in ctx.tables.read_table(acs.CODES_TABLE)] == [
        "0010",
        "0020",
        "0030",
    ]


# ---- mart


def staged_rows(year: int = 2023) -> list[dict[str, object]]:
    parsed = {t: acs.parse_dat(dat(t), year, t) for t in acs.TABLES}
    return acs.build_staged(year, acs.parse_shells(shells(), "B24121"), parsed)


def codes_rows() -> list[dict[str, object]]:
    return acs.build_codes_table(acs.parse_codes(make_xlsx()), acs.parse_crosswalk(make_xlsx()))


def test_build_mart_one_row_per_year_and_occupation_with_total_as_0000() -> None:
    out = acs.build_mart([*staged_rows(2024), *staged_rows(2023)], codes_rows())
    assert [(r["year"], r["occupation_code"]) for r in out] == [
        (2023, "0000"),
        (2023, "0010"),
        (2023, "0020"),
        (2023, "0030"),
        (2024, "0000"),
        (2024, "0010"),
        (2024, "0020"),
        (2024, "0030"),
    ]
    assert out[0] == {
        "iso3": "USA",
        "year": 2023,
        "occupation_code": "0000",
        "occupation": "Total",
        "classification": acs.CLASSIFICATION,
        "soc_code": None,
        "earnings_median": 60192.0,
        "earnings_median_male": 64705.0,
        "earnings_median_female": 52458.0,
        "workers": 118787967.0,
        "workers_male": 66737766.0,
        "workers_female": 52050201.0,
        "earnings_median_moe": 57.0,
        "source": "census_acs1",
    }
    assert out[3]["occupation"] == "Legislators"
    assert out[3]["soc_code"] == "11-1031"
    assert out[3]["earnings_median"] == 69328.0
    assert out[3]["earnings_median_female"] == 58702.0
    assert out[3]["workers"] == 11919.0


def test_build_mart_lists_every_label_that_does_not_match_the_code_table() -> None:
    codes = codes_rows()
    codes[2]["occupation"] = "Lawmakers"
    with pytest.raises(ValueError) as ei:
        acs.build_mart(staged_rows(), codes)
    assert "Legislators" in str(ei.value)
    assert "Chief executives" not in str(ei.value)


def test_build_mart_rejects_when_occupation_counts_differ() -> None:
    codes = [*codes_rows(), {**codes_rows()[0], "occupation_code": "0040", "occupation": "Extra"}]
    with pytest.raises(ValueError):
        acs.build_mart(staged_rows(), codes)


def test_build_mart_rejects_duplicate_year() -> None:
    with pytest.raises(ValueError):
        acs.build_mart([*staged_rows(), *staged_rows()], codes_rows())


def test_mart_skips_when_nothing_is_staged() -> None:
    ctx = make_ctx()
    written, skipped = acs.mart(ctx, [], [])
    assert written == []
    assert skipped == [f"{acs.CODES_TABLE}: no staged table"]
    assert ctx.tables.list_tables() == []


def test_mart_skips_missing_years_and_writes_the_rest() -> None:
    ctx = make_ctx()
    ctx.tables.write_table(acs.CODES_TABLE, codes_rows())
    ctx.tables.write_table(acs.staged_table(2023), staged_rows(2023))
    written, skipped = acs.mart(ctx, [], [])
    assert written == [acs.MART_TABLE]
    assert skipped == [f"{acs.staged_table(y)}: no staged table" for y in (2021, 2022, 2024)]
    assert {r["year"] for r in ctx.tables.read_table(acs.MART_TABLE)} == {2023}


def test_fetch_stage_mart_end_to_end() -> None:
    ctx = make_ctx(responses())
    acs.fetch(ctx, [], [])
    acs.stage(ctx, [], [])
    written, skipped = acs.mart(ctx, [], [])
    assert (written, skipped) == ([acs.MART_TABLE], [])
    rows = ctx.tables.read_table(acs.MART_TABLE)
    assert len(rows) == 4 * 4
    assert {r["source"] for r in rows} == {"census_acs1"}
    assert {r["iso3"] for r in rows} == {"USA"}
