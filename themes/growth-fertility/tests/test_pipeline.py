import json
from pathlib import Path

import pytest

from socioscope_core.core.pipeline import Context, Stage
from socioscope_core.testing.fakes import FakeFetcher, InMemoryRawStore, InMemoryTableStore
from theme_growth_fertility import (
    dhs,
    estat,
    hdr,
    kostat,
    moekr,
    oecd,
    oecd_eag,
    pipeline,
    surveys,
    unsdg,
)
from theme_growth_fertility import eurobarometer as ebm
from theme_growth_fertility import eurostat as es
from theme_growth_fertility import worldbank as wb
from theme_growth_fertility.pipeline import build_panel
from theme_growth_fertility.wiring import PIPELINE

FIXTURES = Path(__file__).parent / "fixtures"
FIXTURE = (FIXTURES / "tfr_page.json").read_bytes()
COUNTRIES = (FIXTURES / "countries_page.json").read_bytes()
DHS_DATA = (FIXTURES / "dhs_tfr_wealth_page.json").read_bytes()
DHS_COUNTRIES = (FIXTURES / "dhs_countries_page.json").read_bytes()
OECD_CSV = (FIXTURES / "oecd_socx_family.csv").read_bytes()
EUROSTAT_FIXTURE_BY_DATASET = {
    "cens_21me_r2": (FIXTURES / "eurostat_cens_fi.json").read_bytes(),
    "demo_fordagec": (FIXTURES / "eurostat_fordagec_fi.json").read_bytes(),
    "demo_pjan": (FIXTURES / "eurostat_pjan_fi.json").read_bytes(),
    "demo_faeduc": (FIXTURES / "eurostat_faeduc_fi.json").read_bytes(),
    "lfsa_pgaed": (FIXTURES / "eurostat_lfsa_fi.json").read_bytes(),
}
# H5 fetch items: one OECD CSV + one Eurostat request per dataset x country
H5_REQUESTS = 1 + len(es.requests())
# H6 items: two PDFs + one Standard Eurobarometer wave each (JSON-LD then Volume A workbook)
H6_ITEMS = 2 + len(ebm.WAVES)
# H8 items: one DHS request per indicator group x breakdown
H8_ITEMS = len(dhs.H8_REQUESTS)
# H8c items: three OECD Education at a Glance CSVs
H8C_ITEMS = len(oecd_eag.DATASETS)
# H8d items: two UN SDG series + one UNDP HDR CSV
H8D_ITEMS = len(unsdg.SERIES) + 1
# H8d-K items: one 교육부 사교육비조사 press-release PDF
H8D_K_ITEMS = 1
# H9 items: one 국가데이터처 출생통계 PDF + one Eurostat demo_frate request per ORDER_GEOS country
H9_ITEMS = 1 + len(es.ORDER_GEOS)


def make_ctx(responses: dict[str, bytes]) -> Context:
    return Context(
        fetcher=FakeFetcher(responses), raw=InMemoryRawStore(), tables=InMemoryTableStore()
    )


def all_responses() -> dict[str, bytes]:
    out = {wb.indicator_url(code): FIXTURE for code in wb.INDICATORS.values()}
    out[wb.countries_url()] = COUNTRIES
    out.update(dhs_responses())
    out.update(h5_responses())
    return out


def h5_responses() -> dict[str, bytes]:
    """OECD CSV + the five Eurostat datasets for Finland only (other countries stay skipped)."""
    out = {oecd.family_spending_url(): OECD_CSV}
    for r in es.requests():
        if r.geo == "FI":
            out[r.url] = EUROSTAT_FIXTURE_BY_DATASET[r.dataset]
    return out


def dhs_responses() -> dict[str, bytes]:
    return {dhs.data_url(): DHS_DATA, dhs.countries_url(): DHS_COUNTRIES}


def test_fetch_stores_raw_with_manifest_and_reports_failures() -> None:
    ctx = make_ctx({wb.indicator_url(wb.INDICATORS["tfr"]): FIXTURE})
    result = PIPELINE.run(Stage.FETCH, ctx)

    assert result.written == ("growth-fertility/worldbank_wdi/tfr.json",)
    # every other indicator + countries + all e-Stat tables + 2 DHS files had no fake response
    assert (
        len(result.skipped)
        == len(wb.INDICATORS)
        + len(estat.TABLES)
        + 2
        + H5_REQUESTS
        + len(kostat.RELEASES)
        + H6_ITEMS
        + H8_ITEMS
        + H8C_ITEMS
        + H8D_ITEMS
        + H8D_K_ITEMS
        + H9_ITEMS
    )
    recs = ctx.raw.records()
    assert len(recs) == 1 and recs[0].license.startswith("CC BY 4.0")
    assert ctx.raw.get(theme="growth-fertility", source="worldbank_wdi", name="tfr.json") == FIXTURE


def test_fetch_rejects_non_json_list_payloads_and_raises_when_all_bad() -> None:
    html = b"<!DOCTYPE html><html><body>parked</body></html>"
    assert wb.is_json_list_payload(FIXTURE) and wb.is_json_list_payload(COUNTRIES)
    assert not wb.is_json_list_payload(html)
    assert not wb.is_json_list_payload(b'{"message": [{"id": "120"}]}')
    ctx = make_ctx({**all_responses(), wb.indicator_url(wb.INDICATORS["tfr"]): html})
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert "tfr: non-JSON-list payload (site down?)" in result.skipped
    assert "growth-fertility/worldbank_wdi/tfr.json" not in result.written
    assert ctx.raw.get(theme="growth-fertility", source="worldbank_wdi", name="tfr.json") is None
    with pytest.raises(RuntimeError, match="worldbank_wdi: all"):
        PIPELINE.run(Stage.FETCH, make_ctx(dict.fromkeys(all_responses(), html)))


def test_fetch_also_stores_country_metadata() -> None:
    ctx = make_ctx(all_responses())
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert "growth-fertility/worldbank_wdi/countries.json" in result.written
    estat_keys = tuple(t.key for t in estat.TABLES)
    h6 = (surveys.TESTA_RAW, surveys.BIB_RAW, "eb_STD")
    assert all(
        s.startswith(
            (
                *estat_keys,
                "eurostat_",
                "newlywed_",
                *h6,
                "dhs_h8_",
                "eag_",
                "sl_dom_",
                "ic_gen_",
                "hdr25_",
                moekr.RAW_NAME,
                kostat.BIRTHS_RAW_NAME,
            )
        )
        for s in result.skipped
    )


def test_stage_then_mart_build_panel_from_raw() -> None:
    ctx = make_ctx(all_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    staged = PIPELINE.run(Stage.STAGE, ctx)
    assert set(staged.written) == {f"staged/worldbank/{k}" for k in wb.INDICATORS} | {
        "staged/worldbank/countries",
        pipeline.DHS_TABLE,
        pipeline.DHS_H8_TABLE,  # the H4 wealth-quintile URL doubles as one H8 request
        pipeline.OECD_TABLE,
        *pipeline.EU_TABLES.values(),
    }
    # the 'NAC' aggregate is dropped at stage using the country metadata
    assert [r["iso3"] for r in ctx.tables.read_table("staged/worldbank/tfr")] == [
        "JPN",
        "JPN",
        "USA",
    ]

    mart = PIPELINE.run(Stage.MART, ctx)
    assert mart.written == (
        "marts/growth_fertility_panel",
        pipeline.DHS_MART,
        pipeline.OECD_MART,
        pipeline.EU_CENSUS_MART,
        pipeline.EU_TFR_ORDER_MART,
        pipeline.EU_TFR_EDU_MART,
        pipeline.DHS_H8_MART,
    )
    panel = ctx.tables.read_table("marts/growth_fertility_panel")
    assert panel[0] == {
        "iso3": "JPN",
        "country": "Japan",
        "year": 2000,
        "region": "East Asia & Pacific",
        "income_group": "High income",
        "tfr": 1.36,
        "gdp_pcap_ppp": 1.36,
        "gdp_growth": 1.36,
        "population": 1.36,
        "u5_mortality": 1.36,
        "fem_sec_enrol": 1.36,
        "urban_share": 1.36,
        "fem_lfp": 1.36,
        "life_exp": 1.36,
        "gdp_pcap_const": 1.36,
        "fem_prim_enrol": 1.36,
        "infant_mortality": 1.36,
        "smam_fe": 1.36,
        "wanted_tfr_wdi": 1.36,
        "married_by18": 1.36,
        "contraceptive_any": 1.36,
        "tert_enrol": 1.36,
        "school_life_exp": 1.36,
        "educ_exp_gdp": 1.36,
        "spend_per_student_sec": 1.36,
        "spend_per_student_tert": 1.36,
        "private_sec_share": 1.36,
    }
    assert [p["iso3"] for p in panel] == ["JPN", "JPN", "USA"]


def test_stage_without_raw_skips_explicitly() -> None:
    ctx = make_ctx({})
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == ()
    assert (
        len(result.skipped)
        == len(wb.INDICATORS)
        + 1
        + len(estat.TABLES)
        + 2
        + H5_REQUESTS
        + len(kostat.RELEASES)
        + H6_ITEMS
        + 1  # H8: the countries raw is missing -> one fail-closed message
        + H8C_ITEMS  # H8c: three raw missing
        + H8D_ITEMS  # H8d: two UN SDG + one HDR raw missing
        + H8D_K_ITEMS  # H8d-K: raw missing
        + H9_ITEMS  # H9: one births PDF + 13 demo_frate raws missing
    )


# ---------------------------------------------------------------- e-Stat (国民生活基礎調査)

ESTAT_FIXTURE_BY_KIND = {
    estat.TableKind.INCOME_DIST_TS: (
        FIXTURES / "estat_children_hh_income_dist_ts.csv"
    ).read_bytes(),
    estat.TableKind.WORKERS_MARITAL: (FIXTURES / "estat_workers_marital_income.csv").read_bytes(),
    estat.TableKind.HH_TYPE: (FIXTURES / "estat_hh_type_income.csv").read_bytes(),
    estat.TableKind.SHUGYO_MARITAL_AGE_INCOME: (
        FIXTURES / "estat_shugyo_marital_age_income.xlsx"
    ).read_bytes(),
}


def estat_responses(tables: tuple[estat.Table, ...] = estat.TABLES) -> dict[str, bytes]:
    return {t.url: ESTAT_FIXTURE_BY_KIND[t.kind] for t in tables}


def test_fetch_stores_estat_files_with_license() -> None:
    ctx = make_ctx(estat_responses())
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert "growth-fertility/estat_kiso/workers_marital_income_2025.csv" in result.written
    assert "growth-fertility/estat_shugyo/shugyo_marital_age_income_2022.xlsx" in result.written
    assert len([w for w in result.written if "/estat_kiso/" in w]) == len(estat.TABLES) - 1
    # no World Bank / DHS / OECD / Eurostat responses
    assert (
        len(result.skipped)
        == len(wb.INDICATORS)
        + 1
        + 2
        + H5_REQUESTS
        + len(kostat.RELEASES)
        + H6_ITEMS
        + H8_ITEMS
        + H8C_ITEMS
        + H8D_ITEMS
        + H8D_K_ITEMS
        + H9_ITEMS
    )
    recs = [r for r in ctx.raw.records() if r.source in {"estat_kiso", "estat_shugyo"}]
    assert len(recs) == len(estat.TABLES)
    assert all("政府標準利用規約" in r.license and "CC BY 4.0" in r.license for r in recs)
    assert all("appId" not in r.url for r in recs)


def test_stage_estat_concatenates_waves_and_reports_missing_ones() -> None:
    present = tuple(t for t in estat.TABLES if t.survey_year in (2025, 2013))
    ctx = make_ctx(estat_responses(present))
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert set(result.written) == set(pipeline.ESTAT_TABLES.values()) - {
        "staged/estat/shugyo_marital_age_income"
    }
    # kiso 2016/19/22 × 2 kinds + the shugyo 2022 table (its key also contains "income_")
    assert sum("raw missing" in s for s in result.skipped if "income_" in s) == 7
    assert "staged/estat/shugyo_marital_age_income" not in result.written
    marital = ctx.tables.read_table("staged/estat/kiso_workers_marital_income")
    assert sorted({int(str(r["survey_year"])) for r in marital}) == [2013, 2025]
    assert {r["source"] for r in marital} == {"estat_kiso"}
    ts = ctx.tables.read_table("staged/estat/kiso_income_dist_ts")
    assert {r["population"] for r in ts} == {"all", "with_children"}


def test_stage_estat_independent_of_worldbank_fail_closed() -> None:
    # No World Bank raw at all: WB stages nothing, e-Stat still stages.
    ctx = make_ctx(estat_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert set(result.written) == set(pipeline.ESTAT_TABLES.values())


def test_mart_jp_income_class_fertility_from_staged() -> None:
    ctx = make_ctx(estat_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    PIPELINE.run(Stage.STAGE, ctx)
    result = PIPELINE.run(Stage.MART, ctx)
    assert result.written == (pipeline.JP_MART, pipeline.JP_AGE_MART)  # no WB -> no WB panel
    rows = ctx.tables.read_table(pipeline.JP_MART)
    assert set(rows[0]) == {
        "survey",
        "survey_year",
        "year",
        "income_class",
        "income_class_lower_yen",
        "income_class_upper_yen",
        "sex",
        "metric",
        "value",
        "denominator",
        "source",
    }
    by = {(r["metric"], r["survey_year"], r["sex"], r["income_class"]): r for r in rows}
    m = by[("married_share", 2025, "male", "50-100")]
    assert m["value"] == pytest.approx(286 / (286 + 582)) and m["denominator"] == 286 + 582
    assert m["income_class_lower_yen"] == 500_000 and m["year"] == 2024
    c = by[("children_household_share", 2013, None, "0-50")]
    assert c["value"] == pytest.approx(2 / 129) and c["denominator"] == 129
    # the stage fixture is the same CSV for both ts tables, so both populations exist
    assert ("children_household_share_pct", 2025, None, "2000-") in by
    assert ("household_share_pct", 2025, None, "2000-") in by


def test_mart_jp_income_age_marital_from_staged_shugyo_only() -> None:
    shugyo = tuple(t for t in estat.TABLES if t.kind is estat.TableKind.SHUGYO_MARITAL_AGE_INCOME)
    ctx = make_ctx(estat_responses(shugyo))
    PIPELINE.run(Stage.FETCH, ctx)
    staged = PIPELINE.run(Stage.STAGE, ctx)
    assert staged.written == ("staged/estat/shugyo_marital_age_income",)
    result = PIPELINE.run(Stage.MART, ctx)
    assert result.written == (pipeline.JP_AGE_MART,)  # kiso not staged -> no kiso mart
    rows = ctx.tables.read_table(pipeline.JP_AGE_MART)
    assert set(rows[0]) == {
        "survey",
        "survey_year",
        "year",
        "sex",
        "age_class",
        "age_lower",
        "age_upper",
        "income_class",
        "income_class_lower_yen",
        "income_class_upper_yen",
        "metric",
        "value",
        "denominator",
        "source",
    }
    assert len(rows) == 2 * 3 * 4  # sex x age x income cells (fixture)
    assert {r["metric"] for r in rows} == {"ever_married_share"}
    assert {r["survey"] for r in rows} == {"就業構造基本調査"}
    by = {(r["sex"], r["age_class"], r["income_class"]): r for r in rows}
    m = by[("male", "25-29", "total")]
    assert m["value"] == pytest.approx((2_928_900 - 2_200_100) / 2_928_900)
    assert m["denominator"] == 2_928_900 and m["year"] == 2022 and m["age_upper"] == 30
    assert by[("male", "85-", "1500-")]["value"] == pytest.approx(1.0)  # never-married '-' = 0
    assert by[("male", "total", "total")]["age_lower"] is None
    # sorted: year, sex, age (total first), income lower (total first)
    assert [r["age_class"] for r in rows[:4]] == ["total"] * 4
    assert rows[0]["income_class"] == "total" and rows[1]["income_class"] == "0-50"


def test_build_income_age_marital_keeps_missing_as_none() -> None:
    base = {
        "survey_year": 2022,
        "year": 2022,
        "source": "estat_shugyo",
        "sex": "male",
        "age_class": "25-29",
        "age_lower": 25,
        "age_upper": 30,
        "income_class": "0-50",
        "income_class_lower_yen": 0,
        "income_class_upper_yen": 500_000,
    }
    mid = {
        "income_class": "50-100",
        "income_class_lower_yen": 500_000,
        "income_class_upper_yen": 1_000_000,
    }
    high = {
        "income_class": "100-150",
        "income_class_lower_yen": 1_000_000,
        "income_class_upper_yen": 1_500_000,
    }
    rows = pipeline.build_income_age_marital(
        [
            {**base, "marital": "total", "persons": 100.0},
            {**base, "marital": "never_married", "persons": None},
            {**base, **mid, "marital": "total", "persons": 0.0},
            {**base, **mid, "marital": "never_married", "persons": 0.0},
            {**base, **high, "marital": "total", "persons": 200.0},
            {**base, **high, "marital": "never_married", "persons": 50.0},
        ]
    )
    assert [(r["income_class"], r["value"], r["denominator"]) for r in rows] == [
        ("0-50", None, 100.0),  # missing numerator
        ("50-100", None, 0.0),  # zero denominator
        ("100-150", 0.75, 200.0),
    ]


def test_build_income_class_fertility_keeps_missing_as_none_and_is_sorted() -> None:
    base = {
        "survey_year": 2025,
        "year": 2024,
        "income_class_lower_yen": 0,
        "income_class_upper_yen": 500_000,
        "income_class": "0-50",
    }
    rows = pipeline.build_income_class_fertility(
        income_dist=[],
        workers_marital=[
            {**base, "sex": "total", "marital": "married", "workers_per_100k": None},
            {**base, "sex": "total", "marital": "unmarried", "workers_per_100k": 5.0},
        ],
        hh_type=[
            {**base, "households_per_10k": 0.0, "with_children_per_10k": 0.0},
            {
                **base,
                "income_class": "2000-",
                "income_class_lower_yen": 20_000_000,
                "income_class_upper_yen": None,
                "households_per_10k": 100.0,
                "with_children_per_10k": 25.0,
            },
        ],
    )
    assert [(r["metric"], r["income_class"], r["value"]) for r in rows] == [
        ("children_household_share", "0-50", None),  # zero denominator -> None
        ("children_household_share", "2000-", 0.25),
        ("married_share", "0-50", None),  # missing numerator -> None
    ]
    assert rows == pipeline.build_income_class_fertility(
        income_dist=[],
        workers_marital=[
            {**base, "sex": "total", "marital": "unmarried", "workers_per_100k": 5.0},
            {**base, "sex": "total", "marital": "married", "workers_per_100k": None},
        ],
        hh_type=[
            {
                **base,
                "income_class": "2000-",
                "income_class_lower_yen": 20_000_000,
                "income_class_upper_yen": None,
                "households_per_10k": 100.0,
                "with_children_per_10k": 25.0,
            },
            {**base, "households_per_10k": 0.0, "with_children_per_10k": 0.0},
        ],
    )


def test_stage_without_country_metadata_fails_closed() -> None:
    # Indicators present, countries.json missing: do not stage aggregates silently.
    ctx = make_ctx({wb.indicator_url(code): FIXTURE for code in wb.INDICATORS.values()})
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == ()
    assert any("countries" in s for s in result.skipped)


def test_build_panel_keeps_missing_as_none() -> None:
    staged = {
        "tfr": [{"iso3": "JPN", "country": "Japan", "year": 2000, "value": 1.36}],
        "gdp_growth": [{"iso3": "JPN", "country": "Japan", "year": 2001, "value": 0.4}],
    }
    assert build_panel(staged, countries=[]) == [
        {
            "iso3": "JPN",
            "country": "Japan",
            "year": 2000,
            "region": None,
            "income_group": None,
            "tfr": 1.36,
            "gdp_growth": None,
        },
        {
            "iso3": "JPN",
            "country": "Japan",
            "year": 2001,
            "region": None,
            "income_group": None,
            "tfr": None,
            "gdp_growth": 0.4,
        },
    ]


# ---------------------------------------------------------------- DHS (TFR by wealth quintile)


def test_wdi_stage_indicators_are_registered() -> None:
    assert wb.INDICATORS["u5_mortality"] == "SH.DYN.MORT"
    assert wb.INDICATORS["fem_sec_enrol"] == "SE.SEC.ENRR.FE"
    assert wb.INDICATORS["urban_share"] == "SP.URB.TOTL.IN.ZS"
    assert wb.INDICATORS["fem_lfp"] == "SL.TLF.CACT.FE.ZS"
    assert wb.INDICATORS["life_exp"] == "SP.DYN.LE00.IN"


def test_wdi_h7_indicators_are_registered() -> None:
    assert wb.INDICATORS["gdp_pcap_const"] == "NY.GDP.PCAP.KD"
    assert wb.INDICATORS["fem_prim_enrol"] == "SE.PRM.ENRR.FE"
    assert wb.INDICATORS["infant_mortality"] == "SP.DYN.IMRT.IN"


def test_stage_writes_h7_indicator_tables_and_mart_has_columns() -> None:
    ctx = make_ctx(all_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    staged = PIPELINE.run(Stage.STAGE, ctx)
    for key in ("gdp_pcap_const", "fem_prim_enrol", "infant_mortality"):
        assert f"staged/worldbank/{key}" in staged.written
        rows = ctx.tables.read_table(f"staged/worldbank/{key}")
        assert set(rows[0]) == {"iso3", "country", "year", "value", "indicator", "source"}
        assert rows[0]["source"] == wb.SOURCE
    PIPELINE.run(Stage.MART, ctx)
    panel = ctx.tables.read_table("marts/growth_fertility_panel")
    assert {"gdp_pcap_const", "fem_prim_enrol", "infant_mortality"} <= set(panel[0])


H8_WDI = {
    "smam_fe": "SP.DYN.SMAM.FE",
    "wanted_tfr_wdi": "SP.DYN.WFRT",
    "married_by18": "SP.M18.2024.FE.ZS",
    "contraceptive_any": "SP.DYN.CONU.ZS",
}


def test_wdi_h8_indicators_are_registered() -> None:
    for key, code in H8_WDI.items():
        assert wb.INDICATORS[key] == code
        assert wb.indicator_url(code) == (
            f"https://api.worldbank.org/v2/country/all/indicator/{code}?format=json&per_page=20000"
        )


def test_stage_writes_h8_indicator_tables_and_mart_has_columns() -> None:
    ctx = make_ctx(all_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    staged = PIPELINE.run(Stage.STAGE, ctx)
    for key, code in H8_WDI.items():
        assert f"staged/worldbank/{key}" in staged.written
        rows = ctx.tables.read_table(f"staged/worldbank/{key}")
        assert set(rows[0]) == {"iso3", "country", "year", "value", "indicator", "source"}
        assert rows[0]["source"] == wb.SOURCE
        assert any(r.url == wb.indicator_url(code) for r in ctx.raw.records())
    PIPELINE.run(Stage.MART, ctx)
    panel = ctx.tables.read_table("marts/growth_fertility_panel")
    assert set(H8_WDI) <= set(panel[0])


H8C_WDI = {
    "tert_enrol": "SE.TER.ENRR",
    "school_life_exp": "SE.SCH.LIFE",
    "educ_exp_gdp": "SE.XPD.TOTL.GD.ZS",
    "spend_per_student_sec": "SE.XPD.SECO.PC.ZS",
    "spend_per_student_tert": "SE.XPD.TERT.PC.ZS",
    "private_sec_share": "SE.SEC.PRIV.ZS",
}


def test_wdi_h8c_indicators_are_registered() -> None:
    for key, code in H8C_WDI.items():
        assert wb.INDICATORS[key] == code
        if key == "school_life_exp":
            continue
        assert wb.indicator_url(code) == (
            f"https://api.worldbank.org/v2/country/all/indicator/{code}?format=json&per_page=20000"
        )


def test_wdi_school_life_exp_url_is_pinned_to_panel_years() -> None:
    # SE.SCH.LIFE is not in WDI (source 2); the API serves it from Education Statistics
    # (source 12) on a 1960-2100 grid (27,472 rows > per_page). Pin the date range so the
    # raw fits one page and the staged grid matches the other indicators (1960-2025).
    assert wb.indicator_url("SE.SCH.LIFE") == (
        "https://api.worldbank.org/v2/country/all/indicator/SE.SCH.LIFE"
        "?format=json&per_page=20000&date=1960:2025"
    )
    # Other indicators are unchanged (their manifest URLs must stay stable).
    assert "date=" not in wb.indicator_url("SP.DYN.TFRT.IN")


def test_stage_writes_h8c_indicator_tables_and_mart_has_columns() -> None:
    ctx = make_ctx(all_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    staged = PIPELINE.run(Stage.STAGE, ctx)
    for key, code in H8C_WDI.items():
        assert f"staged/worldbank/{key}" in staged.written
        rows = ctx.tables.read_table(f"staged/worldbank/{key}")
        assert set(rows[0]) == {"iso3", "country", "year", "value", "indicator", "source"}
        assert rows[0]["source"] == wb.SOURCE
        assert any(r.url == wb.indicator_url(code) for r in ctx.raw.records())
    PIPELINE.run(Stage.MART, ctx)
    panel = ctx.tables.read_table("marts/growth_fertility_panel")
    assert set(H8C_WDI) <= set(panel[0])


def test_fetch_stores_dhs_files_with_citation_license() -> None:
    ctx = make_ctx(dhs_responses())
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert result.written == (
        "growth-fertility/dhs_api/dhs_tfr_wealth.json",
        "growth-fertility/dhs_api/dhs_countries.json",
        # same URL as the H4 wealth-quintile file, stored again under its H8 name
        "growth-fertility/dhs_api/dhs_h8_FE_FRTR_W_TFR_wealth_quintile.json",
    )
    assert (
        len(result.skipped)
        == len(wb.INDICATORS)
        + 1
        + len(estat.TABLES)
        + H5_REQUESTS
        + len(kostat.RELEASES)
        + H6_ITEMS
        + (H8_ITEMS - 1)  # the wealth-quintile request shares the H4 URL
        + H8C_ITEMS
        + H8D_ITEMS
        + H8D_K_ITEMS
        + H9_ITEMS
    )
    recs = ctx.raw.records()
    assert {r.source for r in recs} == {dhs.SOURCE}
    assert all("The DHS Program Indicator Data API" in r.license for r in recs)
    assert {r.url for r in recs} == {dhs.data_url(), dhs.countries_url()}
    assert ctx.raw.get(theme="growth-fertility", source="dhs_api", name="dhs_countries.json") == (
        DHS_COUNTRIES
    )


def test_stage_dhs_writes_long_table_and_reports_unmapped_codes() -> None:
    ctx = make_ctx(dhs_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    # WB / e-Stat absent -> independent; the shared wealth URL also feeds the H8 table
    assert result.written == (pipeline.DHS_TABLE, pipeline.DHS_H8_TABLE)
    assert any("OS" in s and "iso3" in s.lower() for s in result.skipped)
    rows = ctx.tables.read_table(pipeline.DHS_TABLE)
    assert len(rows) == 7
    assert {r["iso3"] for r in rows} == {"AFG", "ALB"}
    assert {r["source"] for r in rows} == {"dhs_api"}


def test_stage_dhs_without_country_metadata_fails_closed() -> None:
    ctx = make_ctx({dhs.data_url(): DHS_DATA})
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == ()
    assert any("dhs_countries" in s and "raw missing" in s for s in result.skipped)


def test_mart_dhs_from_staged_only() -> None:
    ctx = make_ctx(dhs_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    PIPELINE.run(Stage.STAGE, ctx)
    result = PIPELINE.run(Stage.MART, ctx)
    assert result.written == (pipeline.DHS_MART, pipeline.DHS_H8_MART)
    rows = ctx.tables.read_table(pipeline.DHS_MART)
    assert set(rows[0]) == {
        "iso3",
        "country",
        "dhs_country_code",
        "survey_id",
        "survey_year",
        "survey_type",
        "quintile",
        "quintile_label",
        "value",
        "ci_low",
        "ci_high",
        "denominator_weighted",
        "source",
    }
    assert [(r["survey_year"], r["iso3"], r["quintile"]) for r in rows] == [
        (2008, "ALB", 1),
        (2008, "ALB", 5),
        (2015, "AFG", 1),
        (2015, "AFG", 2),
        (2015, "AFG", 3),
        (2015, "AFG", 4),
        (2015, "AFG", 5),
    ]


def test_build_dhs_mart_sorts_and_keeps_rows_unchanged() -> None:
    a = {"survey_year": 2015, "iso3": "AFG", "quintile": 2, "value": 5.4}
    b = {"survey_year": 2008, "iso3": "ALB", "quintile": 5, "value": None}
    c = {"survey_year": 2015, "iso3": "AFG", "quintile": 1, "value": 5.3}
    assert pipeline.build_dhs_mart([a, b, c]) == [b, c, a]


# ---------------------------------------------------------------- H5: OECD SOCX / Eurostat


def test_fetch_stores_oecd_and_eurostat_raw_with_licenses() -> None:
    ctx = make_ctx(h5_responses())
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert result.written[0] == "growth-fertility/oecd_socx/socx_family_spending.csv"
    assert "growth-fertility/eurostat/eurostat_cens_21me_r2_FI.json" in result.written
    assert "growth-fertility/eurostat/eurostat_lfsa_pgaed_FI.json" in result.written
    assert len(result.written) == 1 + 5
    # everything else (WB, e-Stat, DHS, Eurostat for the other countries) is skipped, not fatal
    assert (
        len(result.skipped)
        == len(wb.INDICATORS)
        + 1
        + len(estat.TABLES)
        + 2
        + (H5_REQUESTS - 6)
        + len(kostat.RELEASES)
        + H6_ITEMS
        + H8_ITEMS
        + H8C_ITEMS
        + H8D_ITEMS
        + H8D_K_ITEMS
        + H9_ITEMS
    )
    assert sum(s.startswith("eurostat_cens_21me_r2_") for s in result.skipped) == 30
    recs = ctx.raw.records()
    assert {r.source for r in recs} == {oecd.SOURCE, es.SOURCE}
    oecd_rec = next(r for r in recs if r.source == oecd.SOURCE)
    assert "OECD Terms" in oecd_rec.license and oecd_rec.url == oecd.family_spending_url()
    eu_rec = next(r for r in recs if r.source == es.SOURCE)
    assert "copyright-notice" in eu_rec.license and "cens_21me_r2" in eu_rec.url


def test_stage_oecd_family_spending_independent_of_other_sources() -> None:
    ctx = make_ctx({oecd.family_spending_url(): OECD_CSV})
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == (pipeline.OECD_TABLE,)
    rows = ctx.tables.read_table(pipeline.OECD_TABLE)
    assert {r["iso3"] for r in rows} == {"FIN", "KOR"}
    assert set(rows[0]) == {"iso3", "year", "spending_type", "value", "unit", "source"}
    assert {r["source"] for r in rows} == {"oecd_socx"}
    mart = PIPELINE.run(Stage.MART, ctx)
    assert mart.written == (pipeline.OECD_MART,)
    assert ctx.tables.read_table(pipeline.OECD_MART) == rows  # staged already sorted


def test_stage_eurostat_concatenates_countries_and_reports_missing() -> None:
    ctx = make_ctx(h5_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert set(result.written) == {pipeline.OECD_TABLE, *pipeline.EU_TABLES.values()}
    h5 = [s for s in result.skipped if s.startswith("eurostat_") and "demo_frate" not in s]
    assert sum("raw missing" in s for s in h5) == len(es.requests()) - 5
    census = ctx.tables.read_table(pipeline.EU_TABLES["cens_21me_r2"])
    assert len(census) == 60 and {r["iso3"] for r in census} == {"FIN"}
    assert {r["year"] for r in census} == {2021}
    lfs = ctx.tables.read_table(pipeline.EU_TABLES["lfsa_pgaed"])
    assert len(lfs) == 27 and {r["source"] for r in lfs} == {"eurostat"}


def test_stage_eurostat_without_any_raw_writes_nothing() -> None:
    ctx = make_ctx({})
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == ()
    assert "socx_family_spending.csv: raw missing (run fetch)" in result.skipped


def test_mart_eurostat_tables_from_staged_only() -> None:
    ctx = make_ctx(
        {r.url: EUROSTAT_FIXTURE_BY_DATASET[r.dataset] for r in es.requests() if r.geo == "FI"}
    )
    PIPELINE.run(Stage.FETCH, ctx)
    PIPELINE.run(Stage.STAGE, ctx)
    result = PIPELINE.run(Stage.MART, ctx)
    assert result.written == (
        pipeline.EU_CENSUS_MART,
        pipeline.EU_TFR_ORDER_MART,
        pipeline.EU_TFR_EDU_MART,
    )
    assert any(pipeline.OECD_TABLE in s for s in result.skipped)
    census = ctx.tables.read_table(pipeline.EU_CENSUS_MART)
    assert len(census) == 12 and set(census[0]) == {
        "iso3", "year", "sex", "age_class", "age_lower", "age_upper", "isced_group",
        "married", "total", "married_share", "source",
    }  # fmt: skip
    order = ctx.tables.read_table(pipeline.EU_TFR_ORDER_MART)
    assert [r["order"] for r in order] == ["1", "2", "3", "GE4", "TOTAL", "UNK"]
    assert set(order[0]) == {"iso3", "year", "order", "tfr", "births_age_unknown", "source"}
    edu = ctx.tables.read_table(pipeline.EU_TFR_EDU_MART)
    assert [r["isced_group"] for r in edu] == ["ED0-2", "ED3_4", "ED5-8", "TOTAL"]
    assert set(edu[0]) == {"iso3", "year", "isced_group", "tfr", "women_total_thousand", "source"}


def test_mart_tfr_by_order_needs_both_births_and_women() -> None:
    births_only = next(r for r in es.requests() if r.dataset == "demo_fordagec" and r.geo == "FI")
    ctx = make_ctx({births_only.url: EUROSTAT_FIXTURE_BY_DATASET["demo_fordagec"]})
    PIPELINE.run(Stage.FETCH, ctx)
    PIPELINE.run(Stage.STAGE, ctx)
    result = PIPELINE.run(Stage.MART, ctx)
    assert result.written == ()
    assert any(pipeline.EU_TABLES["demo_pjan"] in s for s in result.skipped)


# ---- H5: Korea newlywed statistics (PDF) ------------------------------------------------


def kostat_release(year: int) -> kostat.Release:
    return next(r for r in kostat.RELEASES if r.ref_year == year)


def kostat_pages(year: int) -> list[str]:
    return (FIXTURES / f"kostat_newlywed_{year}.txt").read_text(encoding="utf-8").split("\f")


def test_fetch_kostat_stores_each_release_pdf_with_kogl_license(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    pdf = pdf_from_pages(kostat_pages(2024))
    ctx = make_ctx({kostat_release(2024).url: pdf, kostat_release(2015).url: b"%PDF-1.4 junk"})
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert "growth-fertility/kostat_newlywed/newlywed_2024.pdf" in result.written
    assert "growth-fertility/kostat_newlywed/newlywed_2015.pdf" in result.written
    assert (
        ctx.raw.get(theme="growth-fertility", source=kostat.SOURCE, name="newlywed_2024.pdf") == pdf
    )
    assert any(s.startswith("newlywed_2016.pdf:") for s in result.skipped)
    recs = [r for r in ctx.raw.records() if r.source == kostat.SOURCE]
    assert len(recs) == 2 and all("KOGL" in r.license for r in recs)
    assert {r.url for r in recs} == {kostat_release(2024).url, kostat_release(2015).url}


def test_stage_kostat_concatenates_releases_and_skips_unrecognised_layout(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    ctx = make_ctx(
        {
            kostat_release(2024).url: pdf_from_pages(kostat_pages(2024)),
            kostat_release(2015).url: pdf_from_pages(kostat_pages(2015)),
            kostat_release(2019).url: pdf_from_pages(["- 1 -\n신혼부부 수\n전체 1,000"]),
            kostat_release(2018).url: b"<html>maintenance</html>",
        }
    )
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert pipeline.KOSTAT_TABLE in result.written
    rows = ctx.tables.read_table(pipeline.KOSTAT_TABLE)
    assert set(rows[0]) == {
        "release_year",
        "ref_year",
        "population",
        "income_class",
        "income_lower_10k_krw",
        "income_upper_10k_krw",
        "couples",
        "with_children_share",
        "children_1_share",
        "children_2_share",
        "children_3plus_share",
        "mean_children",
        "income_concept",
        "source",
    }
    assert sorted({(r["release_year"], r["ref_year"]) for r in rows}) == [
        (2015, 2015),
        (2024, 2023),
        (2024, 2024),
    ]
    assert len(rows) == 21
    assert any(
        s.startswith("newlywed_2019.pdf: layout not recognised (") and "not found" in s
        for s in result.skipped
    )
    assert any(s.startswith("newlywed_2018.pdf: layout not recognised (") for s in result.skipped)
    assert any(s.startswith("newlywed_2016.pdf: raw missing") for s in result.skipped)
    # other sources untouched: nothing of WB / e-Stat / DHS was written
    assert all(name == pipeline.KOSTAT_TABLE for name in result.written)


def test_stage_kostat_writes_nothing_without_any_raw() -> None:
    ctx = make_ctx({})
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert pipeline.KOSTAT_TABLE not in result.written
    assert pipeline.KOSTAT_TABLE not in ctx.tables.list_tables()
    assert sum(s.startswith("newlywed_") and "raw missing" in s for s in result.skipped) == len(
        kostat.RELEASES
    )


def test_mart_kostat_long_table_prefers_latest_release_per_ref_year(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    ctx = make_ctx(
        {
            kostat_release(2024).url: pdf_from_pages(kostat_pages(2024)),
            kostat_release(2023).url: pdf_from_pages(kostat_pages(2023)),  # 2022 + 2023 (dup 2023)
            kostat_release(2015).url: pdf_from_pages(kostat_pages(2015)),  # wage_only
            kostat_release(2016).url: pdf_from_pages(kostat_pages(2016)),  # 2015 earned + 2016
        }
    )
    PIPELINE.run(Stage.FETCH, ctx)
    PIPELINE.run(Stage.STAGE, ctx)
    result = PIPELINE.run(Stage.MART, ctx)
    assert result.written == (pipeline.KOSTAT_MART,)
    rows = ctx.tables.read_table(pipeline.KOSTAT_MART)
    assert set(rows[0]) == {
        "ref_year",
        "release_year",
        "income_class",
        "income_lower_10k_krw",
        "income_upper_10k_krw",
        "metric",
        "value",
        "income_concept",
        "source",
    }
    metrics = {
        "with_children_share",
        "mean_children",
        "couples",
        "children_1_share",
        "children_2_share",
        "children_3plus_share",
    }
    assert {str(r["metric"]) for r in rows} == metrics
    # ref years: 2015 (wage_only + earned_business), 2016, 2022, 2023, 2024 -> 6 blocks × 7 × 6
    assert len(rows) == 6 * 7 * 6
    keyed = {(r["ref_year"], r["income_concept"], r["income_class"], r["metric"]): r for r in rows}
    r2023 = keyed[(2023, "earned_business", "total", "with_children_share")]
    assert r2023["release_year"] == 2024 and r2023["value"] == pytest.approx(0.525)
    assert keyed[(2015, "wage_only", "total", "couples")]["value"] == 852_618
    assert keyed[(2015, "earned_business", "total", "couples")]["value"] == 1_179_000
    assert keyed[(2024, "earned_business", "ge_10000", "mean_children")]["value"] == 0.53
    assert keyed[(2024, "earned_business", "lt_1000", "couples")]["income_upper_10k_krw"] == 1000
    # sorted by ref_year, income_concept, then income lower bound (total first)
    order = [
        (r["ref_year"], r["income_concept"], r["income_class"])
        for r in rows
        if r["metric"] == "couples"
    ]
    assert order[:8] == [
        (2015, "earned_business", "total"),
        (2015, "earned_business", "lt_1000"),
        (2015, "earned_business", "1000_3000"),
        (2015, "earned_business", "3000_5000"),
        (2015, "earned_business", "5000_7000"),
        (2015, "earned_business", "7000_10000"),
        (2015, "earned_business", "ge_10000"),
        (2015, "wage_only", "total"),
    ]
    assert order[-1] == (2024, "earned_business", "ge_10000")


def test_mart_kostat_skipped_when_staged_missing() -> None:
    ctx = make_ctx({})
    result = PIPELINE.run(Stage.MART, ctx)
    assert pipeline.KOSTAT_MART not in result.written
    assert any(
        s.startswith(f"{pipeline.KOSTAT_TABLE}: staged table missing") for s in result.skipped
    )


def test_build_kostat_mart_is_pure_and_deterministic() -> None:
    staged = [
        {
            "release_year": 2021,
            "ref_year": 2021,
            "population": "first_marriage_within_5y",
            "income_class": "total",
            "income_lower_10k_krw": None,
            "income_upper_10k_krw": None,
            "couples": 10,
            "with_children_share": 0.5,
            "children_1_share": 0.4,
            "children_2_share": 0.1,
            "children_3plus_share": 0.0,
            "mean_children": 0.6,
            "income_concept": "earned_business",
            "source": "kostat_newlywed",
        },
        {
            "release_year": 2022,
            "ref_year": 2021,
            "population": "first_marriage_within_5y",
            "income_class": "total",
            "income_lower_10k_krw": None,
            "income_upper_10k_krw": None,
            "couples": 11,
            "with_children_share": 0.5,
            "children_1_share": 0.4,
            "children_2_share": 0.1,
            "children_3plus_share": 0.0,
            "mean_children": 0.6,
            "income_concept": "earned_business",
            "source": "kostat_newlywed",
        },
    ]
    out = pipeline.build_kostat_mart(staged)
    assert len(out) == 6
    assert {r["release_year"] for r in out} == {2022}
    assert next(r for r in out if r["metric"] == "couples")["value"] == 11
    assert out == pipeline.build_kostat_mart(list(reversed(staged)))


# ---------------------------------------------------------------- H6: ideals / expectations surveys

TESTA_PAGES = (FIXTURES / "testa2012_appendix.txt").read_text(encoding="utf-8").split("\f")
BIB_PAGES = (FIXTURES / "bib2025_table1.txt").read_text(encoding="utf-8").split("\f")
STD101 = next(w for w in ebm.WAVES if w.code == "STD101")
VOL_A_URL = (
    "https://webgate.ec.europa.eu/ebsm/api/public/odp/download?key=202289A0E60356141402B692C20F5193"
)


def jsonld_with_vol_a(filename: str, url: str) -> bytes:
    graph = [
        {"@id": "ds", "@type": "dcat:Dataset"},
        {
            "@id": "d1",
            "@type": "dcat:Distribution",
            "dct:title": [{"@language": "en", "@value": f"Link to {filename}"}],
            "dcat:accessURL": {"@id": url},
        },
        {
            "@id": "d2",
            "@type": "dcat:Distribution",
            "dct:title": [{"@language": "en", "@value": "Link to STD101_VOL_AA.xlsx"}],
            "dcat:accessURL": {"@id": "https://x/aa"},
        },
    ]
    return json.dumps({"@graph": graph, "@context": {}}).encode()


def vol_a_xlsx() -> bytes:
    """Minimal VOL A workbook: the life-in-general item for BE / DE / FI (+ aggregates dropped)."""
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QA2_1"
    rows: list[list[object]] = [
        ["Eurobarometer - 101.3"],
        ["VOL A weighted", "Terrain/Fieldwork :  02/04 - 09/05/2024"],
        ["QA2.1. Quelles ...", "QA2.1. What are your expectations for the next twelve months ...?"],
        ["Votre vie en général ", "Your life in general"],
        [],
        [],
        [],
        [],
        ["<<Back to content", "UE27\nEU27", "BE", "D-W", "DE", "D-E", "FI"],
        ["", "Total", 26423, 1006, 1260, 1559, 299, 1001],
        [],
        ["Meilleurs", 1, 1, 1, 1, 1, 1],
        ["Better", 0.26, 0.3, 0.16, 0.17, 0.19, 0.35],
        ["Moins bons", 1, 1, 1, 1, 1, 1],
        ["Worse", 0.17, 0.1, 0.2, 0.23, 0.2, 0.12],
        ["Sans changement", 1, 1, 1, 1, 1, 1],
        ["The same", 0.55, 0.6, 0.6, 0.59, 0.6, 0.5],
        ["Ne sait pas", 1, 1, 1, 1, 1, 1],
        ["Don't know", 0.02, "-", 0.04, 0.01, 0.01, 0.03],
    ]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def h6_responses(pdf_from_pages) -> dict[str, bytes]:  # type: ignore[no-untyped-def]
    return {
        surveys.TESTA_URL: pdf_from_pages(TESTA_PAGES),
        surveys.BIB_URL: pdf_from_pages(BIB_PAGES),
        STD101.jsonld_url: jsonld_with_vol_a("STD101_VOL_A.xlsx", VOL_A_URL),
        VOL_A_URL: vol_a_xlsx(),
    }


def test_fetch_h6_stores_pdfs_and_resolves_volume_a_from_jsonld(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    ctx = make_ctx(h6_responses(pdf_from_pages))
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert "growth-fertility/testa2012_eb75_4/testa2012_edrp_2012_02.pdf" in result.written
    assert "growth-fertility/bib2025_ggs2/bib2025_ggs2_fertility_gaps.pdf" in result.written
    assert "growth-fertility/eurobarometer_std/eb_STD101_meta.json" in result.written
    assert "growth-fertility/eurobarometer_std/eb_STD101_vol_a.xlsx" in result.written
    assert len(result.written) == 4
    # the other 14 waves: one skipped entry each (JSON-LD missing -> workbook not attempted;
    # STD105 withdrawn upstream -> not even requested)
    assert sum(s.startswith("eb_STD") and "meta" in s for s in result.skipped) == len(ebm.WAVES) - 1
    assert VOL_A_URL in ctx.fetcher.requested  # type: ignore[attr-defined]
    recs = {r.name: r for r in ctx.raw.records()}
    assert "not redistributed" in recs["testa2012_edrp_2012_02.pdf"].license
    assert "CC BY-SA 4.0" in recs["bib2025_ggs2_fertility_gaps.pdf"].license
    assert "2011/833/EU" in recs["eb_STD101_vol_a.xlsx"].license
    assert recs["eb_STD101_vol_a.xlsx"].url == VOL_A_URL
    assert recs["eb_STD101_meta.json"].url == STD101.jsonld_url
    assert recs["eb_STD101_meta.json"].source == ebm.SOURCE


def test_fetch_h6_does_not_request_a_withdrawn_wave_and_says_why() -> None:
    std105 = next(w for w in ebm.WAVES if w.code == "STD105")
    ctx = make_ctx({std105.jsonld_url: jsonld_with_vol_a("STD105_VOL_A.xlsx", VOL_A_URL)})
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert std105.jsonld_url not in ctx.fetcher.requested  # type: ignore[attr-defined]
    assert result.written == ()
    (entry,) = [s for s in result.skipped if s.startswith("eb_STD105_meta")]
    assert "withdrawn" in entry and "committed raw" in entry


def test_stage_h6_stages_a_withdrawn_wave_from_its_raw(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    std105 = next(w for w in ebm.WAVES if w.code == "STD105")
    ctx = make_ctx({})  # nothing fetchable: both committed raws are already in the store
    for name, url, payload in (
        (
            std105.meta_raw_name,
            std105.jsonld_url,
            jsonld_with_vol_a("STD105_VOL_A.xlsx", VOL_A_URL),
        ),
        ("eb_STD105_vol_a.xlsx", VOL_A_URL, vol_a_xlsx()),
    ):
        ctx.raw.put(
            theme=pipeline.THEME, source=ebm.SOURCE, name=name, url=url,
            license=ebm.LICENSE, payload=payload,
        )  # fmt: skip
    PIPELINE.run(Stage.STAGE, ctx)
    rows = ctx.tables.read_table(pipeline.EB_EXPECT_TABLE)
    assert {r["wave"] for r in rows} == {std105.code}
    assert all(r["fieldwork_start"] == "2026-03" for r in rows)


def test_fetch_h6_skips_workbook_when_jsonld_is_unusable() -> None:
    ctx = make_ctx({STD101.jsonld_url: b"<html>maintenance</html>"})
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert result.written == ("growth-fertility/eurobarometer_std/eb_STD101_meta.json",)
    assert any(s.startswith("eb_STD101_vol_a") and "Volume A" in s for s in result.skipped)
    assert all("odp/download" not in u for u in ctx.fetcher.requested)  # type: ignore[attr-defined]


def test_stage_h6_writes_ideals_tables_and_expectations(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    ctx = make_ctx(h6_responses(pdf_from_pages))
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert set(result.written) == {
        pipeline.SURVEYS_EB_TABLE,
        pipeline.SURVEYS_GGS_TABLE,
        pipeline.EB_EXPECT_TABLE,
    }
    eb_rows = ctx.tables.read_table(pipeline.SURVEYS_EB_TABLE)
    expected, issues = surveys.eb2011_rows_from_pages(TESTA_PAGES)
    assert eb_rows == expected and len(issues) == 4
    assert sum("garbled" in s for s in result.skipped) == 4
    ggs = ctx.tables.read_table(pipeline.SURVEYS_GGS_TABLE)
    assert len(ggs) == 11 * 4 * 6 and {r["source"] for r in ggs} == {"bib2025_ggs2"}
    exp = ctx.tables.read_table(pipeline.EB_EXPECT_TABLE)
    assert [(r["iso3"], r["better_share"]) for r in exp] == [
        ("BEL", 30.0),
        ("DEU", 17.0),
        ("FIN", 35.0),
    ]
    assert {r["wave"] for r in exp} == {"STD101"} and {r["item"] for r in exp} == {"life_general"}
    assert set(exp[0]) == {
        "iso3", "wave", "fieldwork_year", "fieldwork_half", "fieldwork_start", "item",
        "better_share", "worse_share", "same_share", "dk_share", "source",
    }  # fmt: skip
    assert sum(s.startswith("eb_STD") and "raw missing" in s for s in result.skipped) == (
        len(ebm.WAVES) - 1
    )


def test_stage_h6_unreadable_pdf_or_workbook_is_skipped_not_fatal(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    responses = h6_responses(pdf_from_pages)
    responses[surveys.TESTA_URL] = b"%PDF-1.4 junk"
    responses[VOL_A_URL] = b"<html>maintenance</html>"
    ctx = make_ctx(responses)
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == (pipeline.SURVEYS_GGS_TABLE,)
    assert any(
        s.startswith(surveys.TESTA_RAW) and "layout not recognised" in s for s in result.skipped
    )
    assert any(
        s.startswith("eb_STD101_vol_a") and "layout not recognised" in s for s in result.skipped
    )


def test_mart_h6_ideals_long_table_and_expectations_with_net_optimism(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    ctx = make_ctx(h6_responses(pdf_from_pages))
    PIPELINE.run(Stage.FETCH, ctx)
    PIPELINE.run(Stage.STAGE, ctx)
    result = PIPELINE.run(Stage.MART, ctx)
    assert set(result.written) == {pipeline.IDEALS_MART, pipeline.EXPECT_MART}
    ideals = ctx.tables.read_table(pipeline.IDEALS_MART)
    eb_rows, _ = surveys.eb2011_rows_from_pages(TESTA_PAGES)
    assert ideals == surveys.build_ideals_mart(eb_rows, surveys.ggs_rows_from_pages(BIB_PAGES))
    assert {r["source"] for r in ideals} == {"eb2011", "ggs2020"}
    exp = ctx.tables.read_table(pipeline.EXPECT_MART)
    assert [(r["iso3"], r["net_optimism"]) for r in exp] == [
        ("BEL", 20.0),
        ("DEU", -6.0),
        ("FIN", 23.0),
    ]


def test_mart_h6_ideals_written_from_one_survey_and_skipped_without_any(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    ctx = make_ctx({surveys.BIB_URL: pdf_from_pages(BIB_PAGES)})
    PIPELINE.run(Stage.FETCH, ctx)
    PIPELINE.run(Stage.STAGE, ctx)
    result = PIPELINE.run(Stage.MART, ctx)
    assert result.written == (pipeline.IDEALS_MART,)
    assert {r["source"] for r in ctx.tables.read_table(pipeline.IDEALS_MART)} == {"ggs2020"}
    assert any(s.startswith(pipeline.EB_EXPECT_TABLE) for s in result.skipped)
    empty = make_ctx({})
    result = PIPELINE.run(Stage.MART, empty)
    assert pipeline.IDEALS_MART not in result.written and pipeline.EXPECT_MART not in result.written
    assert any(s.startswith(pipeline.SURVEYS_EB_TABLE) for s in result.skipped)


# ---------------------------------------------------------------- H8: DHS mechanisms (long mart)

H8_COLUMNS = [
    "iso3",
    "country",
    "dhs_country_code",
    "survey_id",
    "survey_year",
    "survey_type",
    "indicator_id",
    "metric",
    "characteristic_category",
    "characteristic_label",
    "by_variable",
    "is_preferred",
    "value",
    "ci_low",
    "ci_high",
    "denominator_weighted",
    "source",
]


def h8_responses() -> dict[str, bytes]:
    return {r.url: (FIXTURES / r.raw_name).read_bytes() for r in dhs.H8_REQUESTS}


def _h8_fixture_rows() -> int:
    return sum(
        len(json.loads((FIXTURES / r.raw_name).read_bytes())["Data"]) for r in dhs.H8_REQUESTS
    )


def test_fetch_h8_stores_every_request_with_citation_license() -> None:
    ctx = make_ctx({**dhs_responses(), **h8_responses()})
    result = PIPELINE.run(Stage.FETCH, ctx)
    expected = tuple(f"growth-fertility/dhs_api/{r.raw_name}" for r in dhs.H8_REQUESTS)
    assert result.written[-len(expected) :] == expected
    recs = {r.name: r for r in ctx.raw.records() if r.name.startswith("dhs_h8_")}
    assert set(recs) == {r.raw_name for r in dhs.H8_REQUESTS}
    for req in dhs.H8_REQUESTS:
        assert recs[req.raw_name].url == req.url
        assert recs[req.raw_name].source == dhs.SOURCE
        assert "The DHS Program Indicator Data API" in recs[req.raw_name].license
    # the H4 raw files are untouched
    assert ctx.raw.get(theme="growth-fertility", source="dhs_api", name="dhs_tfr_wealth.json") == (
        DHS_DATA
    )


def test_stage_h8_writes_long_table_and_reports_unmapped_codes() -> None:
    ctx = make_ctx({**dhs_responses(), **h8_responses()})
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == (pipeline.DHS_TABLE, pipeline.DHS_H8_TABLE)
    assert any("dhs_h8_FE_FRTR_W_TFR_total.json" in s and "OS" in s for s in result.skipped)
    rows = ctx.tables.read_table(pipeline.DHS_H8_TABLE)
    # dropped: OS rows in the TFR total and wealth files, the "Residence" row in the wealth file
    assert len(rows) == _h8_fixture_rows() - 3
    assert list(rows[0]) == H8_COLUMNS
    assert {r["metric"] for r in rows} == set(dhs.METRICS.values())
    assert {r["survey_id"] for r in rows} == {"AF2015DHS", "AL2008DHS", "KE2022DHS"}
    assert {r["source"] for r in rows} == {"dhs_api"}
    assert {r["characteristic_category"] for r in rows} == {
        "Total",
        "Total 15-49",
        "Total 25-49",
        "Education",
        "Wealth quintile",
        "Age (5-year groups)",
    }
    # the H4 staged table is unchanged by the H8 extension
    assert len(ctx.tables.read_table(pipeline.DHS_TABLE)) == 7


def test_stage_h8_missing_raw_is_skipped_and_the_rest_staged() -> None:
    responses = {**dhs_responses(), **h8_responses()}
    missing = dhs.H8_REQUESTS[5]
    del responses[missing.url]
    ctx = make_ctx(responses)
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert pipeline.DHS_H8_TABLE in result.written
    assert any(s.startswith(f"{missing.raw_name}: raw missing") for s in result.skipped)
    rows = ctx.tables.read_table(pipeline.DHS_H8_TABLE)
    assert not any(
        r["metric"] == "wanted_tfr" and r["characteristic_category"] == "Education" for r in rows
    )
    assert any(r["metric"] == "wanted_tfr" for r in rows)


def test_stage_h8_without_country_metadata_fails_closed() -> None:
    ctx = make_ctx(h8_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == ()
    assert any("dhs_countries" in s and "mechanisms" in s for s in result.skipped)


def test_mart_h8_is_sorted_staged_table_with_spec_columns() -> None:
    ctx = make_ctx({**dhs_responses(), **h8_responses()})
    PIPELINE.run(Stage.FETCH, ctx)
    PIPELINE.run(Stage.STAGE, ctx)
    result = PIPELINE.run(Stage.MART, ctx)
    assert result.written == (pipeline.DHS_MART, pipeline.DHS_H8_MART)
    rows = ctx.tables.read_table(pipeline.DHS_H8_MART)
    assert list(rows[0]) == H8_COLUMNS
    assert len(rows) == len(ctx.tables.read_table(pipeline.DHS_H8_TABLE))
    keys = [tuple(r[k] for k in dhs.MECHANISM_KEY) for r in rows]
    assert keys == sorted(keys) and len(set(keys)) == len(keys)
    assert rows[0]["survey_id"] == "AL2008DHS" and rows[-1]["survey_id"] == "KE2022DHS"
    u5 = [r for r in rows if r["metric"] == "u5mr" and r["characteristic_category"] == "Total"]
    assert [(r["survey_id"], r["by_variable"][:4], r["is_preferred"]) for r in u5] == [
        ("AF2015DHS", "Five", 1),
        ("AF2015DHS", "Ten ", 0),
        ("KE2022DHS", "Five", 1),
    ]


def test_mart_h8_duplicate_key_in_staged_fails_closed() -> None:
    ctx = make_ctx({})
    row = {k: "x" for k in H8_COLUMNS}
    row.update(survey_year=2015, value=1.0)
    ctx.tables.write_table(pipeline.DHS_H8_TABLE, [row, dict(row)])
    with pytest.raises(ValueError, match="duplicate"):
        PIPELINE.run(Stage.MART, ctx)


# ---------------------------------------------------------------- H8c: OECD Education at a Glance

EAG_BY_KEY = {
    "fin_source": (FIXTURES / "oecd_eag_fin_source.csv").read_bytes(),
    "earn_rel": (FIXTURES / "oecd_eag_earn_rel.csv").read_bytes(),
    "fin_perstud": (FIXTURES / "oecd_eag_fin_perstud.csv").read_bytes(),
}


def h8c_responses() -> dict[str, bytes]:
    return {d.url: EAG_BY_KEY[d.key] for d in oecd_eag.DATASETS}


def test_fetch_stores_three_eag_raws_with_oecd_license() -> None:
    ctx = make_ctx(h8c_responses())
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert result.written == (
        "growth-fertility/oecd_eag/eag_fin_source.csv",
        "growth-fertility/oecd_eag/eag_earn_rel.csv",
        "growth-fertility/oecd_eag/eag_fin_perstud.csv",
    )
    recs = ctx.raw.records()
    assert {r.source for r in recs} == {oecd_eag.SOURCE}
    assert [r.url for r in recs] == [d.url for d in oecd_eag.DATASETS]
    assert all("OECD Terms" in r.license for r in recs)


def test_stage_and_mart_eag_long_table_independent_of_other_sources() -> None:
    ctx = make_ctx(h8c_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == tuple(pipeline.EAG_TABLES.values())
    assert pipeline.EAG_TABLES == {
        "fin_source": "staged/oecd_eag/fin_source",
        "earn_rel": "staged/oecd_eag/earn_rel",
        "fin_perstud": "staged/oecd_eag/fin_perstud",
    }
    assert "eag_fin_source.csv: non-member REF_AREA dropped: BRA, EU25, OECD" in result.skipped
    assert "eag_earn_rel.csv: non-member REF_AREA dropped: ARG, BRA, OECD" in result.skipped
    fin = ctx.tables.read_table(pipeline.EAG_TABLES["fin_source"])
    assert len(fin) == 6 and {r["dataset"] for r in fin} == {"fin_source"}
    assert {r["source"] for r in fin} == {"oecd_eag"}
    mart = PIPELINE.run(Stage.MART, ctx)
    assert mart.written == (pipeline.EAG_MART,)
    assert pipeline.EAG_MART == "marts/oecd_education_finance"
    rows = ctx.tables.read_table(pipeline.EAG_MART)
    assert len(rows) == 18
    assert {r["dataset"] for r in rows} == {"fin_source", "earn_rel", "fin_perstud"}
    assert {r["iso3"] for r in rows} == {"ITA", "JPN", "KOR", "SWE", "USA"}
    assert set(rows[0]) == {
        "iso3", "year", "dataset", "measure", "education_level", "exp_source", "sex",
        "work_time", "price_base", "unit", "base_per", "value", "obs_status", "source",
    }  # fmt: skip
    keys = [oecd_eag.mart_key(r) for r in rows]
    assert keys == sorted(keys)


def test_stage_eag_with_one_raw_writes_that_table_and_mart_uses_it() -> None:
    ds = oecd_eag.DATASETS[1]
    ctx = make_ctx({ds.url: EAG_BY_KEY[ds.key]})
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == (pipeline.EAG_TABLES["earn_rel"],)
    assert "eag_fin_source.csv: raw missing (run fetch)" in result.skipped
    assert "eag_fin_perstud.csv: raw missing (run fetch)" in result.skipped
    mart = PIPELINE.run(Stage.MART, ctx)
    assert mart.written == (pipeline.EAG_MART,)
    assert len(ctx.tables.read_table(pipeline.EAG_MART)) == 6


def test_mart_eag_skips_without_staged_and_fails_on_duplicate_keys() -> None:
    ctx = make_ctx({})
    result = PIPELINE.run(Stage.MART, ctx)
    assert pipeline.EAG_MART not in result.written
    assert any(s.startswith("staged/oecd_eag/") for s in result.skipped)
    rows = oecd_eag.parse(EAG_BY_KEY["fin_source"], oecd_eag.DATASETS[0]).rows
    ctx.tables.write_table(pipeline.EAG_TABLES["fin_source"], rows + [dict(rows[0])])
    with pytest.raises(ValueError, match="duplicate"):
        PIPELINE.run(Stage.MART, ctx)


# ---------------------------------------------------------------- H8d: UN SDG + UNDP HDR

UNSDG_BY_KEY = {
    "sl_dom_tspd": (FIXTURES / "unsdg_sl_dom_tspd.json").read_bytes(),
    "ic_gen_mgtl": (FIXTURES / "unsdg_ic_gen_mgtl.json").read_bytes(),
}
HDR_CSV = (FIXTURES / "hdr25_composite.csv").read_bytes()


def h8d_responses() -> dict[str, bytes]:
    out = {s.url: UNSDG_BY_KEY[s.key] for s in unsdg.SERIES}
    out[hdr.URL] = HDR_CSV
    return out


def test_fetch_stores_unsdg_and_hdr_raws_with_their_licenses() -> None:
    ctx = make_ctx(h8d_responses())
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert result.written == (
        "growth-fertility/unsdg/sl_dom_tspd.json",
        "growth-fertility/unsdg/ic_gen_mgtl.json",
        "growth-fertility/undp_hdr/hdr25_composite_indices.csv",
    )
    recs = {r.name: r for r in ctx.raw.records()}
    assert recs["sl_dom_tspd.json"].source == "unsdg"
    assert recs["sl_dom_tspd.json"].url == unsdg.SERIES[0].url
    assert "data.un.org" in recs["ic_gen_mgtl.json"].license
    assert "un.org/en/about-us/terms-of-use" in recs["ic_gen_mgtl.json"].license
    assert recs["hdr25_composite_indices.csv"].source == "undp_hdr"
    assert recs["hdr25_composite_indices.csv"].url == hdr.URL
    assert "CC BY 3.0 IGO" in recs["hdr25_composite_indices.csv"].license


def test_stage_and_mart_h8d_tables_independent_of_other_sources() -> None:
    ctx = make_ctx(h8d_responses())
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == (*pipeline.UNSDG_TABLES.values(), pipeline.HDR_TABLE)
    assert pipeline.UNSDG_TABLES == {
        "sl_dom_tspd": "staged/unsdg/sl_dom_tspd",
        "ic_gen_mgtl": "staged/unsdg/ic_gen_mgtl",
    }
    assert pipeline.HDR_TABLE == "staged/hdr/composite"
    assert "sl_dom_tspd.json: aggregate geoArea dropped: 9 Oceania" in result.skipped
    assert "ic_gen_mgtl.json: aggregate geoArea dropped: 9 Oceania" in result.skipped
    assert "sl_dom_tspd.json: 392 Japan 1999 15+ FEMALE: non-numeric value 'N/A'" in result.skipped
    assert "hdr25_composite_indices.csv: aggregate row dropped: ZZK.WORLD World" in result.skipped
    tspd = ctx.tables.read_table(pipeline.UNSDG_TABLES["sl_dom_tspd"])
    assert len(tspd) == 32 and {r["series"] for r in tspd} == {"SL_DOM_TSPD"}
    assert {r["source"] for r in tspd} == {"unsdg"}
    comp = ctx.tables.read_table(pipeline.HDR_TABLE)
    assert {r["source"] for r in comp} == {"undp_hdr"}
    assert {r["iso3"] for r in comp} == {"JPN", "KOR", "SWE", "TUR", "USA"}

    mart = PIPELINE.run(Stage.MART, ctx)
    assert mart.written == (pipeline.UNSDG_MART, pipeline.HDR_MART)
    assert pipeline.UNSDG_MART == "marts/gender_equity"
    assert pipeline.HDR_MART == "marts/hdr_gender"
    rows = ctx.tables.read_table(pipeline.UNSDG_MART)
    assert len(rows) == 40
    assert {r["series"] for r in rows} == {"SL_DOM_TSPD", "IC_GEN_MGTL"}
    assert {r["iso3"] for r in rows} == {"JPN", "KOR", "SWE", "USA"}
    assert set(rows[0]) == {
        "iso3", "country", "year", "series", "sex", "age", "location", "value", "nature",
        "source",
    }  # fmt: skip
    keys = [unsdg.mart_key(r) for r in rows]
    assert keys == sorted(keys) and len(set(keys)) == len(keys)
    hrows = ctx.tables.read_table(pipeline.HDR_MART)
    assert hrows == comp
    assert set(hrows[0]) == {"iso3", "country", "year", "metric", "value", "source"}


def test_stage_h8d_with_one_unsdg_raw_writes_that_table_and_mart_uses_it() -> None:
    s = unsdg.SERIES[1]
    ctx = make_ctx({s.url: UNSDG_BY_KEY[s.key]})
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == (pipeline.UNSDG_TABLES["ic_gen_mgtl"],)
    assert "sl_dom_tspd.json: raw missing (run fetch)" in result.skipped
    assert "hdr25_composite_indices.csv: raw missing (run fetch)" in result.skipped
    mart = PIPELINE.run(Stage.MART, ctx)
    assert mart.written == (pipeline.UNSDG_MART,)
    assert len(ctx.tables.read_table(pipeline.UNSDG_MART)) == 8
    assert f"{pipeline.HDR_TABLE}: staged table missing (run stage)" in mart.skipped


def test_mart_h8d_skips_without_staged_and_fails_on_duplicate_keys() -> None:
    ctx = make_ctx({})
    result = PIPELINE.run(Stage.MART, ctx)
    assert pipeline.UNSDG_MART not in result.written and pipeline.HDR_MART not in result.written
    assert any(s.startswith("staged/unsdg/") for s in result.skipped)
    assert any(s.startswith("staged/hdr/") for s in result.skipped)
    rows = unsdg.parse(UNSDG_BY_KEY["sl_dom_tspd"], unsdg.SERIES[0]).rows
    ctx.tables.write_table(pipeline.UNSDG_TABLES["sl_dom_tspd"], rows + [dict(rows[0])])
    with pytest.raises(ValueError, match="duplicate"):
        PIPELINE.run(Stage.MART, ctx)
    ctx = make_ctx({})
    hrows = hdr.parse(HDR_CSV).rows
    ctx.tables.write_table(pipeline.HDR_TABLE, hrows + [dict(hrows[0])])
    with pytest.raises(ValueError, match="duplicate"):
        PIPELINE.run(Stage.MART, ctx)


# ---- H8d-K: 교육부 사교육비조사 (PDF) -----------------------------------------------------


def moekr_pages() -> list[str]:
    return (FIXTURES / "moe_kr_private_tutoring_2025.txt").read_text(encoding="utf-8").split("\f")


def test_fetch_h8d_k_stores_the_pdf_with_kogl_license(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    pdf = pdf_from_pages(moekr_pages())
    ctx = make_ctx({moekr.URL: pdf})
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert result.written == ("growth-fertility/moe_kr/private_tutoring_2025.pdf",)
    assert ctx.raw.get(theme="growth-fertility", source="moe_kr", name=moekr.RAW_NAME) == pdf
    recs = ctx.raw.records()
    assert len(recs) == 1
    assert recs[0].url == moekr.URL and "KOGL" in recs[0].license and "교육부" in recs[0].license
    assert not any(s.startswith("private_tutoring_2025.pdf") for s in result.skipped)


def test_fetch_h8d_k_reports_fetch_failure_without_writing() -> None:
    ctx = make_ctx({})
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert any(s.startswith("private_tutoring_2025.pdf:") for s in result.skipped)
    assert ctx.raw.get(theme="growth-fertility", source="moe_kr", name=moekr.RAW_NAME) is None


def test_stage_h8d_k_writes_long_table_independently(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    ctx = make_ctx({moekr.URL: pdf_from_pages(moekr_pages())})
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == (pipeline.MOEKR_TABLE,)
    rows = ctx.tables.read_table(pipeline.MOEKR_TABLE)
    assert set(rows[0]) == {"year", "table", "group", "metric", "value", "release_year", "source"}
    assert len(rows) == 85 + 16 + 30
    assert {str(r["table"]) for r in rows} == {"annual", "by_children", "by_level"}
    assert all(r["source"] == "moe_kr" and r["release_year"] == 2026 for r in rows)


def test_stage_h8d_k_skips_unrecognised_layout_and_missing_raw(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    ctx = make_ctx({moekr.URL: pdf_from_pages(["- 1 -\n사교육비 총액\n2025 275,351"])})
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert pipeline.MOEKR_TABLE not in result.written
    assert pipeline.MOEKR_TABLE not in ctx.tables.list_tables()
    assert any(
        s.startswith("private_tutoring_2025.pdf: layout not recognised (") and "annual" in s
        for s in result.skipped
    )
    result = PIPELINE.run(Stage.STAGE, make_ctx({}))
    assert any(s.startswith("private_tutoring_2025.pdf: raw missing") for s in result.skipped)


def test_mart_h8d_k_writes_sorted_long_mart(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    ctx = make_ctx({moekr.URL: pdf_from_pages(moekr_pages())})
    PIPELINE.run(Stage.FETCH, ctx)
    PIPELINE.run(Stage.STAGE, ctx)
    result = PIPELINE.run(Stage.MART, ctx)
    assert result.written == (pipeline.MOEKR_MART,)
    rows = ctx.tables.read_table(pipeline.MOEKR_MART)
    assert len(rows) == 131
    assert (rows[0]["year"], rows[0]["table"], rows[0]["group"], rows[0]["metric"]) == (
        2007,
        "annual",
        "총계",
        "total_100m_krw",
    )
    keyed = {(r["year"], r["table"], r["group"], r["metric"]): r["value"] for r in rows}
    assert len(keyed) == len(rows)
    assert keyed[(2025, "annual", "총계", "spend_per_student_10k_krw")] == 45.8
    assert keyed[(2025, "by_children", "자녀 3명 이상", "participation_rate_pct")] == 67.7
    assert keyed[(2024, "by_level", "고등학교", "spend_per_participant_10k_krw")] == 77.2
    # independent of the other sources: nothing else written, mart missing when staged is missing
    result = PIPELINE.run(Stage.MART, make_ctx({}))
    assert pipeline.MOEKR_MART not in result.written
    assert any(
        s.startswith(f"{pipeline.MOEKR_TABLE}: staged table missing") for s in result.skipped
    )


# ---- H9-KR: 국가데이터처 2025년 출생 통계 (PDF) ------------------------------------------------


def births_pages() -> list[str]:
    return (FIXTURES / "kostat_births_2025.txt").read_text(encoding="utf-8").split("\f")


def test_fetch_h9_kr_stores_the_births_pdf_with_kogl_license(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    pdf = pdf_from_pages(["- 1 -\n2025년 출생 통계"])
    ctx = make_ctx({kostat.BIRTHS_URL: pdf})
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert result.written == ("growth-fertility/kostat_births/births_2025.pdf",)
    assert (
        ctx.raw.get(theme="growth-fertility", source="kostat_births", name="births_2025.pdf") == pdf
    )
    recs = ctx.raw.records()
    assert len(recs) == 1
    assert recs[0].url == kostat.BIRTHS_URL
    assert "KOGL" in recs[0].license and "출생 통계" in recs[0].license
    assert not any(s.startswith("births_2025.pdf") for s in result.skipped)
    result = PIPELINE.run(Stage.FETCH, make_ctx({}))
    assert any(s.startswith("births_2025.pdf:") for s in result.skipped)


# ---- H9-EU: Eurostat demo_frate (single-age fertility rates) -------------------------------


def test_fetch_h9_eu_stores_one_demo_frate_raw_per_country() -> None:
    reqs = es.frate_requests()
    fi = next(r for r in reqs if r.geo == "FI")
    ctx = make_ctx({fi.url: (FIXTURES / "eurostat_frate_fi.json").read_bytes()})
    result = PIPELINE.run(Stage.FETCH, ctx)
    assert result.written == ("growth-fertility/eurostat/eurostat_demo_frate_FI.json",)
    assert sum(s.startswith("eurostat_demo_frate_") for s in result.skipped) == len(reqs) - 1
    recs = ctx.raw.records()
    assert len(recs) == 1 and recs[0].source == es.SOURCE and "copyright-notice" in recs[0].license
    assert recs[0].url == fi.url


def test_stage_and_mart_h9_kr_independent_of_other_sources(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    ctx = make_ctx({kostat.BIRTHS_URL: pdf_from_pages(births_pages())})
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == (pipeline.KR_BIRTHS_TABLE,)
    rows = ctx.tables.read_table(pipeline.KR_BIRTHS_TABLE)
    assert len(rows) == 45 * 17 and set(rows[0]) == {
        "year", "metric", "group", "value", "release_year", "source",
    }  # fmt: skip
    result = PIPELINE.run(Stage.MART, ctx)
    assert result.written == (pipeline.KR_BIRTHS_MART,)
    mart = ctx.tables.read_table(pipeline.KR_BIRTHS_MART)
    assert len(mart) == 45 * 17
    assert (mart[0]["year"], mart[0]["metric"], mart[0]["group"]) == (1981, "births", "계")
    keyed = {(r["year"], r["metric"], r["group"]): r["value"] for r in mart}
    assert keyed[(2025, "births", "첫째아")] == 158684 and keyed[(2025, "tfr", "계")] == 0.8


def test_stage_h9_kr_skips_unrecognised_layout_and_missing_raw(pdf_from_pages) -> None:  # type: ignore[no-untyped-def]
    ctx = make_ctx({kostat.BIRTHS_URL: pdf_from_pages(["- 1 -\n출생아 수\n2025 254,341"])})
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == ()
    assert any(s.startswith("births_2025.pdf: layout not recognised (") for s in result.skipped)
    result = PIPELINE.run(Stage.STAGE, make_ctx({}))
    assert any(s.startswith("births_2025.pdf: raw missing") for s in result.skipped)
    result = PIPELINE.run(Stage.MART, make_ctx({}))
    assert any(
        s.startswith(f"{pipeline.KR_BIRTHS_TABLE}: staged table missing") for s in result.skipped
    )


def test_stage_and_mart_h9_eu_from_the_countries_that_have_raw() -> None:
    fi = next(r for r in es.frate_requests() if r.geo == "FI")
    ctx = make_ctx({fi.url: (FIXTURES / "eurostat_frate_fi.json").read_bytes()})
    PIPELINE.run(Stage.FETCH, ctx)
    result = PIPELINE.run(Stage.STAGE, ctx)
    assert result.written == (pipeline.EU_FRATE_TABLE,)
    missing = [s for s in result.skipped if s.startswith("eurostat_demo_frate_")]
    assert len(missing) == 12 and all("raw missing" in s for s in missing)
    staged = ctx.tables.read_table(pipeline.EU_FRATE_TABLE)
    assert len(staged) == 35 * 35 and {r["iso3"] for r in staged} == {"FIN"}
    assert set(staged[0]) == {"iso3", "year", "age", "rate", "source"}
    result = PIPELINE.run(Stage.MART, ctx)
    assert result.written == (pipeline.EU_COHORT_MART,)
    mart = ctx.tables.read_table(pipeline.EU_COHORT_MART)
    assert set(mart[0]) == {
        "iso3", "metric", "year", "cohort", "age", "value", "censored_at", "source",
    }  # fmt: skip
    assert {str(r["metric"]) for r in mart} == {"asfr", "cohort_cum", "completed"}
    completed = [r for r in mart if r["metric"] == "completed" and r["censored_at"] is None]
    assert [r["cohort"] for r in completed] == list(range(1975, 1980))


def test_stage_h9_eu_without_any_raw_writes_nothing_and_mart_is_skipped() -> None:
    result = PIPELINE.run(Stage.STAGE, make_ctx({}))
    assert pipeline.EU_FRATE_TABLE not in result.written
    result = PIPELINE.run(Stage.MART, make_ctx({}))
    assert any(
        s.startswith(f"{pipeline.EU_FRATE_TABLE}: staged table missing") for s in result.skipped
    )
