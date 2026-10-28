"""Standard Eurobarometer Volume A（data.europa.eu 経由、鍵不要）: 国の議会への信頼。

番外編 R18 の H17c（削減の前後で信頼は変わったか）に使う。Pure transforms only (no I/O)。
波ごとに Volume A（国 × 設問の集計表、xlsx または xls を含む zip）を 1 ファイル取り、
「(NATIONALITY PARLIAMENT) への信頼」の表から「Tend to trust」の行を読む。
値は波によって割合（0.36）か %（36.0）で入っているので % に揃える。EU の集計列は 2019 年の
2 波では「EU28-UK」を EU27 として読む（EU28 も残す）。表が見つからない・値の行が無い・
国列が無い場合は ValueError（fail-closed）。

License: 欧州委員会の再利用方針（Decision 2011/833/EU、出典表示で再利用可。
data.europa.eu のデータセット頁に記載、2026-10-31 確認）。
出典表記「European Commission, Standard Eurobarometer 91–103 (2019–2025), Volume A,
via data.europa.eu」。
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass

SOURCE = "eurobarometer"
LICENSE = (
    "European Commission reuse policy (Decision 2011/833/EU, attribution; https://data.europa.eu)"
)
MAGIC = (b"PK", b"\xd0\xcf\x11\xe0")  # xlsx/zip か xls（OLE2）
METRIC = "trust_national_parliament_pct"


@dataclass(frozen=True)
class Wave:
    number: int
    year: int
    season: str  # spring / summer / autumn / winter
    fieldwork: str  # 調査時期（yyyy-mm）
    url: str
    raw_name: str


_DL = "https://webgate.ec.europa.eu/ebsm/api/public/odp/download?key="
_DL_OLD = "https://webgate.ec.europa.eu/eurobarometer/api/public/odp/download?key="
WAVES: tuple[Wave, ...] = (
    Wave(
        91,
        2019,
        "spring",
        "2019-06",
        _DL_OLD + "BD32911BEED333EC8E89A49A4EE084F5",
        "eb91_volume_A.zip",
    ),
    Wave(
        92,
        2019,
        "autumn",
        "2019-11",
        _DL_OLD + "3A3D4A80BC3648F6E1BF5F2157F58A1E",
        "eb92_volume_A.zip",
    ),
    Wave(
        93, 2020, "summer", "2020-08", _DL + "30C0FB6220D55B50B2DBE9A093B20A11", "eb93_volume_A.zip"
    ),
    Wave(
        94,
        2021,
        "winter",
        "2021-02",
        _DL_OLD + "ACD89D783A5E2C6E54B6B9BA1C5751EF",
        "eb94_volume_A.xlsx",
    ),
    Wave(
        95,
        2021,
        "spring",
        "2021-06",
        _DL + "3DBFC38A29B12076FF0CF68FE03E83CF",
        "eb95_volume_A.xlsx",
    ),
    Wave(
        96,
        2022,
        "winter",
        "2022-01",
        _DL + "380C6267C97A89EEC7117E72F4948595",
        "eb96_volume_A.xlsx",
    ),
    Wave(
        97,
        2022,
        "summer",
        "2022-06",
        _DL + "BFB9C0CCD4CDFA7AE9B5828600C2BCC7",
        "eb97_volume_A.xlsx",
    ),
    Wave(
        98,
        2023,
        "winter",
        "2023-01",
        _DL + "8B1605BC8964BBFFEC49C32E79FF9A63",
        "eb98_volume_A.xlsx",
    ),
    Wave(
        99,
        2023,
        "spring",
        "2023-06",
        _DL + "0CCD7D6A9EDFF30E68F80A00649E91B6",
        "eb99_volume_A.xlsx",
    ),
    Wave(
        100,
        2023,
        "autumn",
        "2023-10",
        _DL + "63ACA1A088544CA14E217226F8E32A1F",
        "eb100_volume_A.xlsx",
    ),
    Wave(
        101,
        2024,
        "spring",
        "2024-04",
        _DL + "202289A0E60356141402B692C20F5193",
        "eb101_volume_A.xlsx",
    ),
    Wave(
        102,
        2024,
        "autumn",
        "2024-10",
        _DL + "A2534D981B6C0CF8878DEB8447FB08F4",
        "eb102_volume_A.xlsx",
    ),
    Wave(
        103,
        2025,
        "spring",
        "2025-05",
        _DL + "AF452C0A6DD70E5D13273C5DF2086956",
        "eb103_volume_A.xlsx",
    ),
)
ITEM = re.compile(r"NATIONALITY PARLIAMENT|NATIONALITY\)\s*Parliament|PARLEMENT NATIONAL", re.I)
QUESTION = re.compile(r"trust|confiance", re.I)
LABEL = re.compile(r"(^|\n)\s*tend to trust\s*$", re.I)  # 設問文の「tend to trust it」は除く
ISO2_TO_ISO3: dict[str, str] = {
    "AT": "AUT",
    "BE": "BEL",
    "BG": "BGR",
    "CY": "CYP",
    "CZ": "CZE",
    "DE": "DEU",
    "DK": "DNK",
    "EE": "EST",
    "EL": "GRC",
    "ES": "ESP",
    "FI": "FIN",
    "FR": "FRA",
    "HR": "HRV",
    "HU": "HUN",
    "IE": "IRL",
    "IT": "ITA",
    "LT": "LTU",
    "LU": "LUX",
    "LV": "LVA",
    "MT": "MLT",
    "NL": "NLD",
    "PL": "POL",
    "PT": "PRT",
    "RO": "ROU",
    "SE": "SWE",
    "SI": "SVN",
    "SK": "SVK",
    "UK": "GBR",
    "EU27": "EU27",
    "EU28": "EU28",
}


def is_magic(payload: bytes) -> bool:
    return payload.startswith(MAGIC)


def _sheets(payload: bytes) -> list[tuple[str, list[list[object]]]]:
    """xlsx / xls / それらを含む zip → [(シート名, 行列)]。"""
    if payload.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(payload)) as z:
            names = z.namelist()
            if "[Content_Types].xml" in names:
                import openpyxl

                wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True)
                return [
                    (name, [list(r) for r in wb[name].iter_rows(values_only=True)])
                    for name in wb.sheetnames
                ]
            inner = [m for m in names if m.lower().endswith((".xls", ".xlsx"))]
            if not inner:
                msg = "eurobarometer: zip without xls/xlsx"
                raise ValueError(msg)
            return _sheets(z.read(sorted(inner)[0]))
    import xlrd

    wb = xlrd.open_workbook(file_contents=payload)
    return [
        (sh.name, [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)])
        for sh in wb.sheets()
    ]


def _code(cell: object) -> str | None:
    """ヘッダのセル → 国コード。2019 年の波の「EU28-UK」は EU27 として扱う。"""
    if not isinstance(cell, str):
        return None
    last = cell.strip().split()[-1].strip() if cell.strip() else ""
    if last == "EU28-UK":
        return "EU27"
    return last if re.fullmatch(r"[A-Z]{2}|EU2[78]", last) else None


def parse_table(rows: Sequence[Sequence[object]]) -> dict[str, float]:
    """議会への信頼の表 → {国コード: tend to trust %}。"""
    hdr = next((r for r in rows if any(_code(c) in ("EU27", "EU28") for c in r)), None)
    if hdr is None:
        msg = "eurobarometer: header row with EU27/EU28 not found"
        raise ValueError(msg)
    idx = next(
        (i for i, r in enumerate(rows) if any(isinstance(c, str) and LABEL.search(c) for c in r)),
        None,
    )
    if idx is None:
        msg = "eurobarometer: 'Tend to trust' row not found"
        raise ValueError(msg)
    cols = [i for i, h in enumerate(hdr) if _code(h)]
    for cand in (rows[idx], rows[idx + 1] if idx + 1 < len(rows) else []):
        vals: list[float] = []
        for i in cols:
            v = cand[i] if i < len(cand) else None
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                vals.append(float(v))
        if vals and max(vals) <= 100:
            scale = 100.0 if max(vals) <= 1 else 1.0
            out: dict[str, float] = {}
            for i in cols:
                v = cand[i] if i < len(cand) else None
                code = _code(hdr[i])
                if code and isinstance(v, (int, float)) and not isinstance(v, bool):
                    out[code] = round(float(v) * scale, 1)
            if "IT" in out:
                return out
    msg = "eurobarometer: no percentage row under 'Tend to trust'"
    raise ValueError(msg)


def parse(payload: bytes, wave: Wave) -> list[dict[str, object]]:
    """Volume A → [{wave, year, season, fieldwork, iso3, metric, value}]（EU 集計を含む）。"""
    if not is_magic(payload):
        msg = "eurobarometer: not an Excel/zip payload"
        raise ValueError(msg)
    for _name, rows in _sheets(payload):
        head = [c for r in rows[:8] for c in r if isinstance(c, str)]
        if not (any(ITEM.search(c) for c in head) and any(QUESTION.search(c) for c in head)):
            continue
        table = parse_table(rows)
        out = [
            {
                "wave": wave.number,
                "year": wave.year,
                "season": wave.season,
                "fieldwork": wave.fieldwork,
                "iso3": ISO2_TO_ISO3[code],
                "metric": METRIC,
                "value": v,
            }
            for code, v in table.items()
            if code in ISO2_TO_ISO3
        ]
        out.sort(key=lambda r: str(r["iso3"]))
        return out
    msg = f"eurobarometer: national parliament trust table not found (wave {wave.number})"
    raise ValueError(msg)


def build_mart(staged: Sequence[Mapping[str, object]]) -> list[dict[str, object]]:
    """国 × 波 の long mart。同じキーが 2 行あれば ValueError。"""
    seen: set[tuple[str, int]] = set()
    out: list[dict[str, object]] = []
    for r in staged:
        key = (str(r["iso3"]), int(str(r["wave"])))
        if key in seen:
            msg = f"eurobarometer: duplicate {key}"
            raise ValueError(msg)
        seen.add(key)
        out.append(
            {
                "iso3": key[0],
                "wave": key[1],
                "year": int(str(r["year"])),
                "season": r["season"],
                "fieldwork": r["fieldwork"],
                "metric": r["metric"],
                "value": r["value"],
                "source": SOURCE,
            }
        )
    out.sort(key=lambda r: (str(r["iso3"]), int(str(r["wave"]))))
    return out
